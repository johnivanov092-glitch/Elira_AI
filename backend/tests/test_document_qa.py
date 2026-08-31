from __future__ import annotations

from unittest.mock import patch

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from app.application.code_agent.document_validation import (
    _pdf_structure_issues,
    infer_expected_page_count,
    validate_document,
)
from app.application.code_agent.tools._content import tool_file_gen
from app.application.code_agent.tools._resources import tool_resource_publish


def test_glued_centered_docx_heading_fails_validation(tmp_path) -> None:
    path = tmp_path / "proposal.docx"
    document = Document()
    heading = document.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading.add_run("ОФИСНЫЙПКДЛЯБухГАЛТЕРА—DDR5").bold = True
    document.add_paragraph("Корректный основной текст документа.")
    document.save(path)

    with (
        patch(
            "app.application.code_agent.document_validation._render_to_pdf",
            return_value=(path, "test"),
        ),
        patch(
            "app.application.code_agent.document_validation._page_count",
            return_value=1,
        ),
        patch(
            "app.application.code_agent.document_validation._inspect_pages",
            return_value=("passed", []),
        ) as inspect_pages,
    ):
        result = validate_document(path, expected_page_count=1)

    assert result["status"] == "failed"
    assert result["sha256"]
    assert result["page_count"] == 1
    inspect_pages.assert_called_once_with(path, 1)
    assert {issue["code"] for issue in result["issues"]} == {
        "suspicious_glued_heading",
    }


def test_glued_heading_in_pdf_text_layer_fails_before_vision(tmp_path) -> None:
    path = tmp_path / "proposal.pdf"
    path.write_bytes(b"pdf-bytes-are-not-parsed-because-reader-is-mocked")
    page = type("Page", (), {
        "extract_text": lambda self: "ОФИСНЫЙПКДЛЯБухГАЛТЕРА—DDR5",
    })()
    reader = type("Reader", (), {"pages": [page]})()

    with patch("pypdf.PdfReader", return_value=reader):
        issues = _pdf_structure_issues(path)

    assert {issue["code"] for issue in issues} == {"suspicious_glued_heading"}


def test_resource_publish_rejects_document_that_failed_external_qa(tmp_path) -> None:
    path = tmp_path / "proposal.docx"
    path.write_bytes(b"not-used-because-validation-is-mocked")
    qa = {
        "status": "failed",
        "sha256": "a" * 64,
        "format": "docx",
        "renderer": "microsoft_word",
        "page_count": 2,
        "expected_page_count": 1,
        "vision_status": "failed",
        "issues": [{"code": "page_count_mismatch", "message": "Expected 1 page, got 2."}],
    }

    with patch(
        "app.application.code_agent.tools._resources.validate_document",
        return_value=qa,
    ):
        result = tool_resource_publish(
            tmp_path,
            project_path="proposal.docx",
            expected_page_count=1,
        )

    assert result["ok"] is False
    assert result["error"] == "document_validation_failed"
    assert result["document_qa"] == {**qa, "target": "proposal.docx", "attempt": 1}


def test_resource_publish_returns_hash_bound_document_qa(tmp_path) -> None:
    path = tmp_path / "proposal.docx"
    path.write_bytes(b"not-used-because-validation-is-mocked")
    sha256 = "b" * 64
    qa = {
        "status": "passed",
        "sha256": sha256,
        "format": "docx",
        "renderer": "microsoft_word",
        "page_count": 1,
        "expected_page_count": 1,
        "vision_status": "passed",
        "issues": [],
    }

    with (
        patch(
            "app.application.code_agent.tools._resources.validate_document",
            return_value=qa,
        ),
        patch(
            "app.application.media.resource_store.publish_copy",
            return_value=(123, sha256),
        ),
    ):
        result = tool_resource_publish(
            tmp_path,
            project_path="proposal.docx",
            expected_page_count=1,
        )

    assert result["ok"] is True
    assert result["sha256"] == sha256
    assert result["document_qa"] == {**qa, "target": "proposal.docx"}


def test_resource_publish_stops_after_two_identical_document_qa_failures(tmp_path) -> None:
    path = tmp_path / "proposal.docx"
    path.write_bytes(b"same-invalid-document")
    sha256 = "d" * 64
    qa = {
        "status": "failed",
        "sha256": sha256,
        "format": "docx",
        "renderer": "not_run",
        "page_count": None,
        "expected_page_count": 1,
        "vision_status": "not_run",
        "issues": [{"code": "suspicious_glued_heading", "message": "Broken heading."}],
    }

    with (
        patch(
            "app.application.code_agent.tools._resources.document_sha256",
            return_value=sha256,
        ),
        patch(
            "app.application.code_agent.tools._resources.validate_document",
            return_value=qa,
        ) as validator,
    ):
        first = tool_resource_publish(
            tmp_path,
            project_path="proposal.docx",
            expected_page_count=1,
            run_id="qa-retry-cap",
        )
        second = tool_resource_publish(
            tmp_path,
            project_path="proposal.docx",
            expected_page_count=1,
            run_id="qa-retry-cap",
        )
        third = tool_resource_publish(
            tmp_path,
            project_path="proposal.docx",
            expected_page_count=1,
            run_id="qa-retry-cap",
        )

    assert first["document_qa"]["attempt"] == 1
    assert second["document_qa"]["attempt"] == 2
    assert third["error"] == "document_validation_attempts_exhausted"
    assert validator.call_count == 2


def test_resource_publish_retry_cap_is_per_artifact_even_when_bytes_change(tmp_path) -> None:
    path = tmp_path / "proposal-changing.docx"
    path.write_bytes(b"revision-one")
    qa = {
        "status": "failed",
        "sha256": "a" * 64,
        "format": "docx",
        "renderer": "not_run",
        "page_count": None,
        "expected_page_count": 1,
        "vision_status": "not_run",
        "issues": [{"code": "layout_issue", "message": "Broken layout."}],
    }

    with (
        patch(
            "app.application.code_agent.tools._resources.document_sha256",
            side_effect=["a" * 64, "b" * 64, "c" * 64],
        ),
        patch(
            "app.application.code_agent.tools._resources.validate_document",
            return_value=qa,
        ) as validator,
    ):
        first = tool_resource_publish(
            tmp_path,
            project_path=path.name,
            expected_page_count=1,
            run_id="qa-artifact-retry-cap",
        )
        path.write_bytes(b"revision-two")
        second = tool_resource_publish(
            tmp_path,
            project_path=path.name,
            expected_page_count=1,
            run_id="qa-artifact-retry-cap",
        )
        path.write_bytes(b"revision-three")
        third = tool_resource_publish(
            tmp_path,
            project_path=path.name,
            expected_page_count=1,
            run_id="qa-artifact-retry-cap",
        )

    assert first["document_qa"]["attempt"] == 1
    assert second["document_qa"]["attempt"] == 2
    assert third["error"] == "document_validation_attempts_exhausted"
    assert validator.call_count == 2


def test_file_gen_does_not_return_download_for_unverified_document(tmp_path) -> None:
    generated = tmp_path / "generated.docx"
    generated.write_bytes(b"generated-document")
    qa = {
        "status": "unverified",
        "sha256": "e" * 64,
        "format": "docx",
        "renderer": "unavailable",
        "page_count": None,
        "expected_page_count": 1,
        "vision_status": "not_run",
        "issues": [{"code": "renderer_unavailable", "message": "No renderer."}],
    }

    with (
        patch(
            "app.application.skills.generate_word",
            return_value={
                "ok": True,
                "path": str(generated),
                "filename": "generated.docx",
                "size": generated.stat().st_size,
                "download_url": "/api/skills/download/generated.docx",
            },
        ),
        patch(
            "app.application.code_agent.tools._content.validate_document",
            return_value=qa,
        ),
    ):
        result = tool_file_gen(
            tmp_path,
            format="word",
            title="Report",
            content="Body",
            expected_page_count=1,
        )

    assert result["ok"] is False
    assert result["error"] == "document_validation_unverified"
    assert result["document_qa"] == {
        **qa,
        "attempt": 1,
        "target": result["download_name"],
    }
    assert "download_url" not in result


def test_document_tools_reject_fractional_page_count_before_work(tmp_path) -> None:
    path = tmp_path / "proposal.docx"
    path.write_bytes(b"document")

    publish = tool_resource_publish(
        tmp_path,
        project_path="proposal.docx",
        expected_page_count=1.5,  # type: ignore[arg-type]
    )
    with patch("app.application.skills.generate_word") as generator:
        generated = tool_file_gen(
            tmp_path,
            format="word",
            expected_page_count=1.5,  # type: ignore[arg-type]
        )

    assert publish["error"] == "invalid_expected_page_count"
    assert generated["error"] == "invalid_expected_page_count"
    generator.assert_not_called()


def test_page_count_contract_is_inferred_only_from_explicit_document_wording() -> None:
    assert infer_expected_page_count("Оба КП: 1 страница на документ") == 1
    assert infer_expected_page_count("Сделай одностраничный PDF") == 1
    assert infer_expected_page_count("Документ должен быть ровно 2 страницы") == 2
    assert infer_expected_page_count("Проверь страницу 2 сайта") is None
