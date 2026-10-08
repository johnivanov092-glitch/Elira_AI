"""Standalone mutable document skills: real format roundtrips and honest gaps."""
from __future__ import annotations

import importlib.util
import json
from pathlib import Path
import subprocess
import sys
import zipfile

import pytest


SKILLS = Path(__file__).resolve().parents[2] / "skills"


@pytest.mark.parametrize("name", ["document-read", "document-create", "ocr"])
def test_skill_frontmatter_is_accepted_by_real_catalog_parser(name):
    from app.application.code_agent.task_skills import read_package
    result = read_package(name, SKILLS / name)
    assert result["name"] == name and result["description"] and result["content"]


def load_script(skill: str, filename: str):
    spec = importlib.util.spec_from_file_location(f"standalone_{skill.replace('-', '_')}", SKILLS / skill / filename)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def run_script(skill: str, script: str, *args):
    result = subprocess.run([sys.executable, str(SKILLS / skill / script), *map(str, args)],
                            capture_output=True, text=True, encoding="utf-8", timeout=30)
    return result.returncode, json.loads(result.stdout)


@pytest.mark.parametrize("kind", ["docx", "xlsx", "pdf"])
def test_generated_documents_roundtrip_russian_numbers_and_tail(tmp_path, kind):
    pytest.importorskip({"docx": "docx", "xlsx": "openpyxl", "pdf": "reportlab"}[kind])
    reader = load_script("document-read", "read_document.py")
    payload = {"title": "Отчёт", "content": "## Проверка\nПривет, мир!\n- Конец записи 12345",
               "headers": ["Имя", "Сумма"], "data": [["Привет, мир!", 12345], ["Конец записи", 9.5]],
               "tables": [[["Товар", "Число"], ["Последняя строка", 17]]]}
    source = tmp_path / "input.json"
    source.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
    output = tmp_path / f"result.{kind}"
    code, generated = run_script("document-create", "create_document.py", "--format", kind,
                                 "--input", source, "--output", output)
    assert code == 0, generated
    result = reader.read_document(output)
    assert result["complete"] is True, result
    assert "Привет, мир!" in result["text"] and "Конец записи" in result["text"] and "12345" in result["text"]
    if kind == "docx":
        assert "Последняя строка | 17" in result["text"]


def test_stdout_excerpt_does_not_truncate_saved_text(tmp_path):
    source, output = tmp_path / "text.txt", tmp_path / "full.txt"
    source.write_text("начало\n" + "середина\n" * 100 + "КОНЕЦ", encoding="utf-8")
    code, result = run_script("document-read", "read_document.py", source, "--output", output, "--max-chars", 12)
    assert code == 0 and result["truncated"] is True and result["next_offset"] == 12
    assert output.read_text(encoding="utf-8") == source.read_text(encoding="utf-8")
    assert output.read_text(encoding="utf-8").endswith("КОНЕЦ")


def test_scan_pdf_reports_ocr_required_without_running_network(tmp_path):
    pytest.importorskip("reportlab")
    from PIL import Image, ImageDraw
    from reportlab.pdfgen.canvas import Canvas
    reader = load_script("document-read", "read_document.py")
    image = tmp_path / "scan.png"
    scan = Image.new("RGB", (500, 100), "white")
    ImageDraw.Draw(scan).text((20, 20), "SCANNED DOCUMENT 12345", fill="black")
    scan.save(image)
    pdf = tmp_path / "scan.pdf"
    canvas = Canvas(str(pdf))
    canvas.drawImage(str(image), 30, 600, width=500, height=100)
    canvas.save()
    code, result = run_script("document-read", "read_document.py", pdf)
    assert code == 2 and result["ok"] is True and result["complete"] is False
    assert result["ocr_required"] == [1]


def test_zip_reads_text_without_extracting_traversal_paths(tmp_path):
    source = tmp_path / "archive.zip"
    with zipfile.ZipFile(source, "w") as archive:
        archive.writestr("../outside.txt", "Проверенный конец")
        archive.writestr("voice.ogg", b"OggS fake")
    reader = load_script("document-read", "read_document.py")
    result = reader.read_document(source)
    assert "Проверенный конец" in result["text"] and "OggS fake" not in result["text"]
    assert not (tmp_path.parent / "outside.txt").exists()
    assert result["entries"][1]["text_read"] is False


def test_pptx_reads_slide_text_and_tables(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches
    source = tmp_path / "slides.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Первый слайд"
    table = slide.shapes.add_table(1, 2, Inches(1), Inches(2), Inches(5), Inches(1)).table
    table.cell(0, 0).text, table.cell(0, 1).text = "Итог", "12345"
    presentation.slides.add_slide(presentation.slide_layouts[5]).shapes.title.text = "КОНЕЦ ПРЕЗЕНТАЦИИ 56789"
    presentation.save(source)
    result = load_script("document-read", "read_document.py").read_document(source)
    assert "Первый слайд" in result["text"] and "Итог | 12345" in result["text"]
    assert result["text"].endswith("КОНЕЦ ПРЕЗЕНТАЦИИ 56789")


def test_invalid_legacy_xls_is_honest_error(tmp_path):
    source = tmp_path / "bad.xls"
    source.write_bytes(b"not a BIFF workbook")
    code, result = run_script("document-read", "read_document.py", source)
    assert code == 1 and result["ok"] is False


def test_creation_refuses_overwrite_and_excel_formulas_are_explicit(tmp_path):
    from openpyxl import load_workbook
    payload = tmp_path / "input.json"
    payload.write_text('{"data":[["=1+2"]]}', encoding="utf-8")
    output = tmp_path / "data.xlsx"
    code, result = run_script("document-create", "create_document.py", "--format", "xlsx", "--input", payload, "--output", output)
    assert code == 0, result
    before = output.read_bytes()
    code, result = run_script("document-create", "create_document.py", "--format", "xlsx", "--input", payload, "--output", output)
    assert code == 1 and result["ok"] is False and output.read_bytes() == before
    workbook = load_workbook(output, data_only=False)
    assert workbook.active["A1"].value == "=1+2" and workbook.active["A1"].data_type == "s"
    workbook.close()


def test_legacy_binary_doc_is_honest_failure(tmp_path):
    source = tmp_path / "legacy.doc"
    source.write_bytes(b"not OOXML")
    code, result = run_script("document-read", "read_document.py", source)
    assert code == 1 and result["ok"] is False and "Unsupported format .doc" in result["message"]


def test_docx_read_keeps_table_between_surrounding_paragraphs(tmp_path):
    from docx import Document
    source = tmp_path / "order.docx"
    document = Document()
    document.add_paragraph("BEFORE")
    document.add_table(1, 1).cell(0, 0).text = "MIDDLE"
    document.add_paragraph("AFTER")
    document.save(source)
    text = load_script("document-read", "read_document.py").read_document(source)["text"]
    assert text.index("BEFORE") < text.index("MIDDLE") < text.index("AFTER")


def test_legacy_text_encoding_is_explicit_and_never_silently_mojibake(tmp_path):
    source = tmp_path / "legacy.txt"
    source.write_bytes("Привет".encode("cp866"))
    code, result = run_script("document-read", "read_document.py", source)
    assert code == 1 and result["ok"] is False
    code, result = run_script("document-read", "read_document.py", source, "--encoding", "cp866")
    assert code == 0 and result["text"] == "Привет"


def test_ocr_partial_errors_and_invalid_shapes_are_not_complete():
    ocr = load_script("ocr", "ocr.py")
    result = ocr.normalize_response({"pages": [{"page": 1, "text": "Один", "blocks": [{"box": [1, 2, 3, 4]}]},
                                             {"page": 2, "text": "", "errors": [{"code": "failed"}]}], "errors": []})
    assert result["text"] == "Один\n\n" and result["complete"] is False
    assert result["pages"][0]["blocks"][0]["box"] == [1, 2, 3, 4]
    assert len(result["errors"]) == 2
    with pytest.raises(ValueError):
        ocr.normalize_response({"text": "Pretend success", "pages": "invalid"})
    explicit = ocr.normalize_response({"text": "Page 1", "pages": [{"page": 1, "text": "Page 1"}], "complete": False})
    assert explicit["complete"] is False and explicit["errors"][0]["code"] == "server_incomplete"


def test_uncalculated_excel_formula_is_retained_with_honest_incomplete_status(tmp_path):
    from openpyxl import Workbook
    path = tmp_path / "formula.xlsx"
    workbook = Workbook()
    workbook.active["A1"] = "=1+2"
    workbook.save(path)
    workbook.close()
    result = load_script("document-read", "read_document.py").read_document(path)
    assert "=1+2" in result["text"] and result["complete"] is False
    assert result["formula_cells"][0]["cached_value"] is None
    assert result["errors"][0]["code"] == "formula_value_missing"


def test_tiny_sparse_xlsx_with_huge_dimensions_fails_before_row_expansion(tmp_path):
    from openpyxl import Workbook
    source = tmp_path / "sparse.xlsx"
    workbook = Workbook()
    workbook.active["XFD1048576"] = "one actual cell"
    workbook.save(source)
    workbook.close()
    assert source.stat().st_size < 10000
    code, result = run_script("document-read", "read_document.py", source)
    assert code == 1 and result["ok"] is False and "too many cells" in result["message"]


def test_document_url_download_keeps_provenance_and_bounds_bytes(tmp_path, monkeypatch):
    import hashlib
    import requests
    reader = load_script("document-read", "read_document.py")
    class Response:
        def __init__(self, url, code=200, data=b"document", headers=None):
            self.url, self.status_code, self.data, self.headers = url, code, data, headers or {}
        def __enter__(self):
            return self
        def __exit__(self, *args):
            return False
        def raise_for_status(self):
            if self.status_code >= 400:
                raise requests.HTTPError("download failed")
        def iter_content(self, chunk_size):
            yield self.data
    responses = [Response("https://example.org/start", 302, headers={"Location": "/final.pdf"}),
                 Response("https://example.org/final.pdf")]
    seen = []
    def get(url, **kwargs):
        seen.append((url, kwargs))
        return responses.pop(0)
    monkeypatch.setattr(requests, "get", get)
    destination = tmp_path / "original.pdf"
    provenance = reader.download_document("https://example.org/start", destination, 7, 8)
    assert destination.read_bytes() == b"document"
    assert provenance == {"requested_url": "https://example.org/start", "final_url": "https://example.org/final.pdf",
                          "local_path": str(destination.resolve()), "sha256": hashlib.sha256(b"document").hexdigest()}
    assert seen[0][1]["timeout"] == (7, 7) and seen[0][1]["allow_redirects"] is False
    responses.append(Response("https://example.org/large", data=b"too many bytes"))
    with pytest.raises(ValueError, match="max-bytes"):
        reader.download_document("https://example.org/large", tmp_path / "large.pdf", 7, 2)
    assert not (tmp_path / "large.pdf").exists()
    responses.append(Response("https://example.org/start", 302, headers={"Location": "file:///secret"}))
    with pytest.raises(ValueError, match="HTTP"):
        reader.download_document("https://example.org/start", tmp_path / "bad.pdf", 7, 100)
    def timeout(*args, **kwargs):
        raise requests.Timeout("fixture timeout")
    monkeypatch.setattr(requests, "get", timeout)
    with pytest.raises(requests.Timeout):
        reader.download_document("https://example.org/timeout", tmp_path / "timeout.pdf", 7, 100)


def test_ocr_uses_confirmed_multipart_contract_and_rejects_redirects(tmp_path, monkeypatch):
    import requests
    ocr = load_script("ocr", "ocr.py")
    source = tmp_path / "sample.png"
    source.write_bytes(b"synthetic")
    captured = {}
    class Response:
        status_code = 200
        def raise_for_status(self):
            pass
        def json(self):
            return {"text": "Recognized", "pages": [{"page": 1, "text": "Recognized"}], "errors": []}
    response = Response()
    def post(url, **kwargs):
        captured.update(url=url, **kwargs)
        assert kwargs["files"]["file"][1].read() == b"synthetic"
        return response
    monkeypatch.setattr(requests, "post", post)
    assert ocr.request_ocr(source, "http://ocr.local/ocr", "ru", 12, True)["complete"] is True
    assert captured["data"] == {"language": "ru", "pdf_fallback": "true"}
    assert captured["timeout"] == 12 and captured["allow_redirects"] is False
    response.status_code = 307
    with pytest.raises(ValueError, match="redirected"):
        ocr.request_ocr(source, "http://ocr.local/ocr", "ru", 12, True)


def test_pdf_ocr_default_uploads_original_file_and_preserves_server_pages(tmp_path, monkeypatch, capsys):
    import requests
    source = tmp_path / "two-pages.pdf"
    original = b"%PDF-synthetic-original-bytes"
    source.write_bytes(original)
    saved = tmp_path / "ocr.json"
    calls = []
    class Response:
        status_code = 200
        def raise_for_status(self):
            pass
        def json(self):
            return {"text": "Page one\nPage two", "pages": [
                {"page": 1, "text": "Page one"},
                {"page": 2, "text": "Page two", "blocks": [{"bbox": [1, 2, 3, 4]}]},
            ], "errors": []}
    def post(url, **kwargs):
        filename, handle, mime = kwargs["files"]["file"]
        calls.append((url, filename, handle.read(), mime, kwargs["data"]))
        return Response()
    monkeypatch.setattr(requests, "post", post)
    monkeypatch.setattr(sys, "argv", ["ocr.py", str(source), "--url", "http://ocr.local/ocr",
                                     "--json-output", str(saved)])
    ocr = load_script("ocr", "ocr.py")
    assert ocr.main() == 0
    result = json.loads(saved.read_text(encoding="utf-8"))
    stdout = json.loads(capsys.readouterr().out)
    assert calls == [("http://ocr.local/ocr", "two-pages.pdf", original, "application/pdf",
                      {"language": "auto", "pdf_fallback": "true"})]
    assert result["pdf_mode"] == "server" and result["whole_document_complete"] is True
    assert result["total_pages"] == 2 and result["requested_pages"] is None
    assert result["pages"][1]["page"] == 2 and result["pages"][1]["blocks"][0]["bbox"] == [1, 2, 3, 4]
    assert stdout["pdf_mode"] == "server" and stdout["page_count"] == 2


@pytest.mark.parametrize("args", [["--pdf-mode", "images"], ["--pages", "2"],
                                  ["--dpi", "150"], ["--max-pages", "1"]])
def test_pdf_ocr_rejects_client_render_options_before_http(tmp_path, monkeypatch, args):
    import requests
    def forbidden(*arguments, **kwargs):
        pytest.fail("Rejected client rendering must not send HTTP")
    monkeypatch.setattr(requests, "post", forbidden)
    monkeypatch.setattr(sys, "argv", ["ocr.py", str(tmp_path / "sample.pdf"),
                                     "--url", "http://ocr.local/ocr", *args])
    with pytest.raises(SystemExit) as exc:
        load_script("ocr", "ocr.py").main()
    assert exc.value.code == 2


@pytest.mark.parametrize("response", [
    {"text": "Page one", "pages": [{"page": 1, "text": "Page one"}], "page_count": 2},
    {"text": "Page two", "pages": [{"page": 2, "text": "Page two"}]},
    {"text": "Aggregate without pages", "pages": []},
])
def test_pdf_ocr_missing_server_pages_are_not_complete(tmp_path, monkeypatch, response):
    ocr = load_script("ocr", "ocr.py")
    monkeypatch.setattr(ocr, "request_ocr", lambda *args: ocr.normalize_response(response))
    result = ocr.request_pdf_ocr(tmp_path / "sample.pdf", "http://ocr.local/ocr", "auto", 15)
    assert result["complete"] is False and result["whole_document_complete"] is False
    assert any(error["code"] == "incomplete_pdf_pages" for error in result["errors"])


def test_pdf_ocr_transport_failure_has_no_local_fallback_or_retry(tmp_path, monkeypatch, capsys):
    import requests
    source = tmp_path / "sample.pdf"
    source.write_bytes(b"%PDF-fixture")
    calls = []
    def unavailable(*args, **kwargs):
        calls.append(1)
        raise requests.ConnectionError("fixture server unavailable")
    monkeypatch.setattr(requests, "post", unavailable)
    monkeypatch.setattr(sys, "argv", ["ocr.py", str(source), "--url", "http://ocr.local/ocr"])
    assert load_script("ocr", "ocr.py").main() == 1
    result = json.loads(capsys.readouterr().out)
    assert result["ok"] is False and result["error"] == "ConnectionError"
    assert calls == [1]
