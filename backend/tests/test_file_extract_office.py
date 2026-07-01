from __future__ import annotations

import io
import unittest
from unittest.mock import patch

from app.application.file_extract.runtime import extract_file


class OfficeExtractTest(unittest.TestCase):
    """.pptx (python-pptx) and legacy .xls (xlrd) extraction in the composer path."""

    def test_pptx_roundtrip(self) -> None:
        # Build a real .pptx with python-pptx, then extract its text back out.
        from pptx import Presentation
        from pptx.util import Inches

        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout
        box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
        box.text_frame.text = "Привет из презентации"
        buf = io.BytesIO()
        prs.save(buf)

        out = extract_file("deck.pptx", buf.getvalue())
        self.assertTrue(out["ok"])
        self.assertIn("Привет из презентации", out["text"])
        self.assertIn("Слайд 1", out["text"])

    def test_xls_routed_to_xls_extractor_not_xlsx(self) -> None:
        # No xlwt to synthesise a real .xls here, so verify the dispatch: .xls must
        # go to _extract_xls (xlrd) — openpyxl (_extract_xlsx) cannot read .xls.
        with patch(
            "app.application.file_extract.runtime._extract_xls", return_value="XLS_TEXT"
        ) as m_xls, patch(
            "app.application.file_extract.runtime._extract_xlsx", return_value="XLSX_TEXT"
        ) as m_xlsx:
            out = extract_file("legacy.xls", b"\xd0\xcf\x11\xe0 fake-ole")
        m_xls.assert_called_once()
        m_xlsx.assert_not_called()
        self.assertEqual(out["text"], "XLS_TEXT")


if __name__ == "__main__":
    unittest.main()
