from __future__ import annotations

import io

from app.application.library import runtime
from app.application.media.resource_store import ResourceRecord


def _isolated_library(tmp_path, monkeypatch) -> None:
    monkeypatch.setattr(runtime, "SQLITE_DB", tmp_path / "library.db")
    monkeypatch.setattr(runtime, "UPLOADS_DIR", tmp_path / "uploads")
    runtime.init_library_db()


def test_library_keeps_full_text_but_injects_only_relevant_excerpt(
    tmp_path, monkeypatch,
) -> None:
    _isolated_library(tmp_path, monkeypatch)
    body = "начало\n" + ("обычный текст " * 1200) + "DEEP_LIBRARY_TARGET\nконец"

    added = runtime.add_file_contents(
        filename="long-note.txt",
        contents=body.encode("utf-8"),
        content_type="text/plain",
        use_in_context=True,
    )
    context = runtime.build_library_context(
        query="Найди DEEP_LIBRARY_TARGET",
        max_files=1,
        max_chars_per_file=500,
    )

    assert added["content_chars"] == len(body)
    assert added["preview_len"] == 12000
    assert context["used_files"] == ["long-note.txt"]
    assert "DEEP_LIBRARY_TARGET" in context["context"]
    assert len(context["context"]) < 700


def test_unrelated_request_does_not_receive_recent_library_files(
    tmp_path, monkeypatch,
) -> None:
    _isolated_library(tmp_path, monkeypatch)
    runtime.add_file_contents(
        filename="private-note.txt",
        contents="совершенно другая тема".encode("utf-8"),
        use_in_context=True,
    )

    context = runtime.build_library_context(query="Какая сегодня погода?")

    assert context["selection"] == "no_match"
    assert context["used_files"] == []
    assert context["context"] == ""


def test_library_read_pages_through_the_complete_extracted_text(
    tmp_path, monkeypatch,
) -> None:
    _isolated_library(tmp_path, monkeypatch)
    body = "".join(str(index % 10) for index in range(25000))
    added = runtime.add_file_contents(
        filename="book.txt",
        contents=body.encode("utf-8"),
        content_type="text/plain",
        use_in_context=False,
    )

    first = runtime.read_library_file(added["id"], offset=0, limit=8000)
    second = runtime.read_library_file(
        added["id"], offset=first["next_offset"], limit=8000,
    )
    tail = runtime.read_library_file(added["id"], offset=24000, limit=8000)

    assert first["text"] == body[:8000]
    assert first["has_more"] is True
    assert second["text"] == body[8000:16000]
    assert tail["text"] == body[24000:]
    assert tail["has_more"] is False
    assert tail["next_offset"] is None


def test_library_full_text_includes_docx_tables(tmp_path, monkeypatch) -> None:
    from docx import Document

    _isolated_library(tmp_path, monkeypatch)
    document = Document()
    document.add_paragraph("Вводный текст")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Ключ"
    table.cell(0, 1).text = "TABLE_TARGET"
    payload = io.BytesIO()
    document.save(payload)

    added = runtime.add_file_contents(
        filename="table.docx",
        contents=payload.getvalue(),
        use_in_context=False,
    )
    page = runtime.read_library_file(added["id"], limit=10000)

    assert "Вводный текст" in page["text"]
    assert "TABLE_TARGET" in page["text"]


def test_library_full_text_reads_late_xlsx_sheets_and_rows(tmp_path, monkeypatch) -> None:
    from openpyxl import Workbook

    _isolated_library(tmp_path, monkeypatch)
    workbook = Workbook()
    workbook.remove(workbook.active)
    for sheet_index in range(6):
        sheet = workbook.create_sheet(f"Sheet{sheet_index + 1}")
        for row_index in range(250):
            sheet.append([f"row-{row_index}"])
    workbook["Sheet6"]["A250"] = "LATE_XLSX_TARGET"
    payload = io.BytesIO()
    workbook.save(payload)
    workbook.close()

    added = runtime.add_file_contents(
        filename="large.xlsx",
        contents=payload.getvalue(),
        use_in_context=False,
    )
    search = runtime.search_files("LATE_XLSX_TARGET")

    assert search["items"][0]["id"] == added["id"]
    assert "LATE_XLSX_TARGET" in search["items"][0]["excerpt"]


def test_import_resource_reuses_the_existing_library_owner(
    tmp_path, monkeypatch,
) -> None:
    _isolated_library(tmp_path, monkeypatch)
    blob = tmp_path / "resource-blob"
    blob.write_text("полный текст ресурса", encoding="utf-8")
    record = ResourceRecord(
        resource_id="a" * 32,
        original_name="resource.txt",
        kind="document",
        content_type="text/plain",
        size=blob.stat().st_size,
        sha256="b" * 64,
        created_at=1.0,
        owner_session="session",
        storage_path=str(blob),
    )

    monkeypatch.setattr(
        "app.application.media.resource_store.get_record",
        lambda resource_id: record if resource_id == record.resource_id else None,
    )
    monkeypatch.setattr(
        "app.application.media.resource_store.read_bytes",
        lambda resolved: blob.read_bytes() if resolved == record else b"",
    )

    imported = runtime.import_resource(record.resource_id, use_in_context=True)
    listed = runtime.list_library_files()

    assert imported["ok"] is True
    assert imported["name"] == "resource.txt"
    assert imported["active"] is True
    assert listed["count"] == 1
    assert listed["files"][0]["status"] == "ready"
    assert listed["files"][0]["content_chars"] == len("полный текст ресурса")


def test_import_resource_rejects_unknown_id(tmp_path, monkeypatch) -> None:
    _isolated_library(tmp_path, monkeypatch)
    monkeypatch.setattr(
        "app.application.media.resource_store.get_record",
        lambda _resource_id: None,
    )

    result = runtime.import_resource("f" * 32)

    assert result == {"ok": False, "error": "resource_not_found"}


def test_corrupt_document_is_not_reported_as_indexed(tmp_path, monkeypatch) -> None:
    _isolated_library(tmp_path, monkeypatch)

    result = runtime.add_file_contents(
        filename="broken.docx",
        contents=b"not-a-docx",
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )

    assert result["ok"] is True
    assert result["status"] == "failed"
    assert result["content_chars"] == 0


def test_deleting_legacy_duplicate_keeps_shared_file(tmp_path, monkeypatch) -> None:
    _isolated_library(tmp_path, monkeypatch)
    shared = tmp_path / "uploads" / "shared.txt"
    shared.write_text("shared", encoding="utf-8")
    connection = runtime._conn()
    try:
        connection.executemany(
            "INSERT INTO files (name, stored_path, sha256) VALUES (?, ?, ?)",
            [
                ("first.txt", str(shared), "a" * 64),
                ("second.txt", str(shared), "a" * 64),
            ],
        )
        connection.commit()
        first_id = int(connection.execute(
            "SELECT id FROM files WHERE name = 'first.txt'",
        ).fetchone()[0])
    finally:
        connection.close()

    runtime.delete_file(first_id)

    assert shared.is_file()
