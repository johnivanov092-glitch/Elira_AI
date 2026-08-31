"""Server-owned QA for downloadable PDF and DOCX artifacts.

The code model may choose document contents, but it does not decide whether QA
passed. This module binds structural checks, rendered page count, and an
independent vision inspection to the exact SHA-256 later published to the user.
"""
from __future__ import annotations

import base64
import hashlib
import io
import json
import logging
import os
import re
import shutil
import subprocess
import tempfile
from collections import OrderedDict
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from threading import Lock
from typing import Any

logger = logging.getLogger(__name__)

_SUPPORTED_EXTENSIONS = {".docx", ".pdf"}
# Four parallel calls keep an ordinary 1-12 page business document inside the
# 120-second tool budget. Longer documents fail closed instead of being only
# partially inspected or silently blocking the Qwen tool loop for minutes.
_MAX_VISION_PAGES = 12
_VISION_WORKERS = 4
_WORD_TIMEOUT_SECONDS = 45
_VISION_TIMEOUT_SECONDS = 15.0
_GLUED_TOKEN_RE = re.compile(r"[A-Za-zА-Яа-яЁё]{20,}", re.UNICODE)
_GLUED_TITLE_SIGNALS = ("ДЛЯ", "СБОРКА", "КОМПЬЮТЕР", "DDR", "ПК")
DOCUMENT_QA_ATTEMPT_LIMIT = 2
_ATTEMPT_CACHE_LIMIT = 2048
_ATTEMPTS: OrderedDict[tuple[str, str, int | None], int] = OrderedDict()
_ATTEMPTS_LOCK = Lock()

_LAYOUT_PROMPT = """Ты внешний QA-валидатор страницы документа. Проверь только видимую вёрстку:
- обрезанный, наложенный или вышедший за границы текст;
- склеенные слова или пропавшие пробелы, особенно в заголовках;
- разорванные таблицы, скрытые строки, пустые хвостовые страницы;
- нечитаемые символы или сломанные шрифты.
Не оценивай содержание и подбор товаров. Ответь строго одним JSON-объектом без Markdown:
{"layout_issue": false, "issues": []}
Если дефект есть, layout_issue=true, а issues — короткий список конкретных дефектов на русском.
"""


def _issue(code: str, message: str) -> dict[str, str]:
    return {"code": code, "message": message}


def document_sha256(path: str | Path) -> str:
    path = Path(path)
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def normalize_expected_page_count(value: Any) -> int | None:
    """Accept Qwen's common integer/string-integer forms; reject coercive floats."""
    if value is None:
        return None
    if isinstance(value, bool):
        raise ValueError("expected_page_count must be an integer from 1 to 100")
    if isinstance(value, str) and value.strip().isdigit():
        value = int(value)
    elif not isinstance(value, int):
        raise ValueError("expected_page_count must be an integer from 1 to 100")
    if not 1 <= value <= 100:
        raise ValueError("expected_page_count must be an integer from 1 to 100")
    return value


def infer_expected_page_count(task_text: str | None) -> int | None:
    """Extract only explicit document page-count contracts from user text."""
    text = " ".join(str(task_text or "").casefold().split())
    if not text:
        return None
    if re.search(r"\b(?:одностраничн\w*|one[- ]page|single[- ]page)\b", text):
        return 1
    patterns = (
        r"\b(\d{1,3})\s+страниц(?:а|ы)?\s+на\s+(?:кажд\w*\s+)?(?:документ|файл|кп)\b",
        r"\b(?:документ|файл|кп)\w*[^.!?]{0,60}\b(?:ровно|строго|по|в)\s+(\d{1,3})\s+страниц(?:а|ы)?\b",
        r"\b(?:exactly|strictly)\s+(\d{1,3})\s+pages?\b",
    )
    for pattern in patterns:
        match = re.search(pattern, text, re.IGNORECASE | re.UNICODE)
        if match:
            try:
                return normalize_expected_page_count(match.group(1))
            except ValueError:
                return None
    return None


def document_qa_attempts(
    run_id: str,
    artifact_key: str,
    expected_page_count: int | None,
) -> int:
    if not run_id or not artifact_key:
        return 0
    key = (run_id, artifact_key.casefold(), expected_page_count)
    with _ATTEMPTS_LOCK:
        count = _ATTEMPTS.get(key, 0)
        if key in _ATTEMPTS:
            _ATTEMPTS.move_to_end(key)
        return count


def record_document_qa_failure(
    run_id: str,
    artifact_key: str,
    expected_page_count: int | None,
) -> int:
    if not run_id or not artifact_key:
        return 1
    key = (run_id, artifact_key.casefold(), expected_page_count)
    with _ATTEMPTS_LOCK:
        count = _ATTEMPTS.get(key, 0) + 1
        _ATTEMPTS[key] = count
        _ATTEMPTS.move_to_end(key)
        while len(_ATTEMPTS) > _ATTEMPT_CACHE_LIMIT:
            _ATTEMPTS.popitem(last=False)
        return count


def clear_document_qa_failures(run_id: str, artifact_key: str) -> None:
    if not run_id or not artifact_key:
        return
    with _ATTEMPTS_LOCK:
        for key in tuple(_ATTEMPTS):
            if key[0] == run_id and key[1] == artifact_key.casefold():
                _ATTEMPTS.pop(key, None)


def _base_result(path: Path, expected_page_count: int | None) -> dict[str, Any]:
    return {
        "status": "unverified",
        "sha256": document_sha256(path),
        "format": path.suffix.lower().lstrip("."),
        "renderer": "not_run",
        "page_count": None,
        "expected_page_count": expected_page_count,
        "vision_status": "not_run",
        "issues": [],
    }


def _looks_like_glued_heading(text: str) -> bool:
    for token in _GLUED_TOKEN_RE.findall(text):
        upper = token.upper()
        signals = sum(signal in upper[1:-1] for signal in _GLUED_TITLE_SIGNALS)
        mixed_case_break = bool(re.search(r"[А-ЯЁ]{5,}[а-яё]{1,}[А-ЯЁ]", token))
        if signals >= 2 or mixed_case_break:
            return True
    return False


def _docx_structure_issues(path: Path) -> list[dict[str, str]]:
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        document = Document(path)
    except Exception:  # noqa: BLE001 - never expose package/parser internals
        return [_issue("invalid_docx", "DOCX не открывается как корректный документ.")]

    visible_text: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        visible_text.append(text)
        is_heading_like = (
            paragraph.alignment == WD_ALIGN_PARAGRAPH.CENTER
            or any(run.bold is True for run in paragraph.runs)
        )
        if is_heading_like and _looks_like_glued_heading(text):
            return [_issue(
                "suspicious_glued_heading",
                "В заголовке обнаружена подозрительно длинная склейка слов без пробелов.",
            )]

    for table in document.tables:
        for row in table.rows:
            visible_text.extend(cell.text.strip() for cell in row.cells if cell.text.strip())
    if not visible_text:
        return [_issue("empty_document", "Документ не содержит видимого текста.")]
    return []


def _pdf_structure_issues(path: Path) -> list[dict[str, str]]:
    """Catch deterministic text defects before consulting the vision model.

    Text extraction is best-effort because scanned PDFs legitimately have no
    text layer; rendering and vision remain the authoritative fallback there.
    """
    try:
        from pypdf import PdfReader

        pages = PdfReader(str(path)).pages
        extracted = "\n".join((page.extract_text() or "") for page in pages)
    except Exception:  # noqa: BLE001 - render/page-count checks handle invalid PDFs
        return []
    for line in extracted.splitlines():
        candidate = line.strip()
        if candidate and len(candidate) <= 160 and _looks_like_glued_heading(candidate):
            return [_issue(
                "suspicious_glued_heading",
                "В текстовом слое PDF обнаружена подозрительная склейка слов без пробелов.",
            )]
    return []


def _powershell_literal(value: Path) -> str:
    return "'" + str(value).replace("'", "''") + "'"


def _cleanup_owned_word(pid_path: Path) -> None:
    try:
        pid = int(pid_path.read_text(encoding="ascii").strip())
    except (OSError, ValueError):
        return
    if pid <= 0:
        return
    try:
        subprocess.run(
            ["taskkill.exe", "/PID", str(pid), "/T", "/F"],
            check=False,
            capture_output=True,
            timeout=10,
            creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
        )
    except (OSError, subprocess.SubprocessError):
        logger.warning("document QA could not clean up owned Word pid=%s", pid)


def _render_docx_with_word(source: Path, output: Path, pid_path: Path) -> bool:
    powershell = shutil.which("powershell.exe") or shutil.which("pwsh.exe")
    if not powershell or os.name != "nt":
        return False
    script = f"""
$ErrorActionPreference = 'Stop'
$word = $null
$document = $null
$before = @(Get-Process WINWORD -ErrorAction SilentlyContinue | ForEach-Object {{ $_.Id }})
try {{
  $word = New-Object -ComObject Word.Application
  $word.Visible = $false
  $word.DisplayAlerts = 0
  $after = @(Get-Process WINWORD -ErrorAction SilentlyContinue | ForEach-Object {{ $_.Id }})
  $owned = @($after | Where-Object {{ $_ -notin $before }} | Select-Object -First 1)
  if ($owned.Count -gt 0) {{
    [System.IO.File]::WriteAllText({_powershell_literal(pid_path)}, [string]$owned[0], [System.Text.Encoding]::ASCII)
  }}
  $document = $word.Documents.Open({_powershell_literal(source)}, $false, $true)
  $document.ExportAsFixedFormat({_powershell_literal(output)}, 17)
}} finally {{
  if ($document -ne $null) {{ $document.Close($false); [void][Runtime.InteropServices.Marshal]::ReleaseComObject($document) }}
  if ($word -ne $null) {{ $word.Quit(); [void][Runtime.InteropServices.Marshal]::ReleaseComObject($word) }}
  [GC]::Collect()
  [GC]::WaitForPendingFinalizers()
}}
"""
    encoded = base64.b64encode(script.encode("utf-16le")).decode("ascii")
    try:
        completed = subprocess.run(
            [powershell, "-NoProfile", "-NonInteractive", "-EncodedCommand", encoded],
            check=False,
            capture_output=True,
            timeout=_WORD_TIMEOUT_SECONDS,
            creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
        )
    except subprocess.TimeoutExpired:
        _cleanup_owned_word(pid_path)
        return False
    except OSError:
        return False
    finally:
        pid_path.unlink(missing_ok=True)
    return completed.returncode == 0 and output.is_file() and output.stat().st_size > 0


def _render_docx_with_libreoffice(source: Path, output_dir: Path) -> Path | None:
    executable = shutil.which("soffice") or shutil.which("libreoffice")
    if not executable:
        return None
    try:
        completed = subprocess.run(
            [
                executable,
                "--headless",
                "--convert-to",
                "pdf",
                "--outdir",
                str(output_dir),
                str(source),
            ],
            check=False,
            capture_output=True,
            timeout=_WORD_TIMEOUT_SECONDS,
        )
    except (OSError, subprocess.SubprocessError):
        return None
    candidate = output_dir / f"{source.stem}.pdf"
    return candidate if completed.returncode == 0 and candidate.is_file() else None


def _render_to_pdf(source: Path, output_dir: Path) -> tuple[Path | None, str]:
    if source.suffix.lower() == ".pdf":
        return source, "source_pdf"

    word_pdf = output_dir / "rendered.pdf"
    if _render_docx_with_word(source, word_pdf, output_dir / "word.pid"):
        return word_pdf, "microsoft_word"
    libreoffice_pdf = _render_docx_with_libreoffice(source, output_dir)
    if libreoffice_pdf is not None:
        return libreoffice_pdf, "libreoffice"
    return None, "unavailable"


def _page_count(pdf_path: Path) -> int | None:
    try:
        from pypdf import PdfReader

        return len(PdfReader(str(pdf_path)).pages)
    except Exception:  # noqa: BLE001 - stable QA result, no parser details
        return None


def _parse_layout_response(response: str) -> tuple[bool, list[str]] | None:
    match = re.search(r"\{.*\}", response or "", flags=re.DOTALL)
    if not match:
        return None
    try:
        payload = json.loads(match.group(0))
    except (TypeError, ValueError):
        return None
    layout_issue = payload.get("layout_issue")
    issues = payload.get("issues")
    if not isinstance(layout_issue, bool) or not isinstance(issues, list):
        return None
    clean_issues = [str(item).strip()[:300] for item in issues if str(item).strip()]
    return layout_issue, clean_issues


def _inspect_pages(pdf_path: Path, pages: int) -> tuple[str, list[dict[str, str]]]:
    if pages < 1:
        return "failed", [_issue("empty_render", "После рендера в документе нет страниц.")]
    if pages > _MAX_VISION_PAGES:
        return "unverified", [_issue(
            "vision_page_limit",
            f"Vision-QA проверяет не более {_MAX_VISION_PAGES} страниц; в документе их {pages}.",
        )]
    try:
        from pdf2image import convert_from_path
        from app.infrastructure.llm.vision_ocr import describe_image

        images = convert_from_path(str(pdf_path), dpi=120, fmt="png")
    except Exception:  # noqa: BLE001 - dependency/provider details stay in logs
        return "unverified", [_issue("page_render_unavailable", "Не удалось растрировать страницы для vision-QA.")]
    if len(images) != pages:
        return "unverified", [_issue("page_render_incomplete", "Не все страницы переданы на vision-QA.")]

    def inspect_page(item: tuple[int, Any]) -> tuple[int, str, list[dict[str, str]]]:
        index, image = item
        buffer = io.BytesIO()
        image.save(buffer, format="PNG")
        response = describe_image(
            f"page-{index}.png",
            buffer.getvalue(),
            prompt=_LAYOUT_PROMPT,
            timeout_seconds=_VISION_TIMEOUT_SECONDS,
        )
        parsed = _parse_layout_response(response or "")
        if parsed is None:
            return index, "unverified", [_issue(
                "vision_response_invalid",
                f"Vision-QA не вернул структурированный результат для страницы {index}.",
            )]
        layout_issue, page_issues = parsed
        if layout_issue:
            if not page_issues:
                page_issues = ["Vision-валидатор обнаружил дефект вёрстки."]
            return index, "failed", [
                _issue("layout_issue", f"Page {index}: {message}")
                for message in page_issues
            ]
        return index, "passed", []

    with ThreadPoolExecutor(max_workers=min(_VISION_WORKERS, pages)) as executor:
        inspected = list(executor.map(inspect_page, enumerate(images, start=1)))
    inspected.sort(key=lambda item: item[0])
    if any(status == "unverified" for _, status, _ in inspected):
        issues = [issue for _, _, page_issues in inspected for issue in page_issues]
        return "unverified", issues
    issues = [issue for _, _, page_issues in inspected for issue in page_issues]
    return ("failed", issues) if issues else ("passed", [])


def validate_document(
    path: str | Path,
    *,
    expected_page_count: int | None = None,
) -> dict[str, Any]:
    """Validate one immutable document snapshot and return JSON-safe evidence."""
    target = Path(path)
    result = _base_result(target, expected_page_count)
    if target.suffix.lower() not in _SUPPORTED_EXTENSIONS:
        result["issues"] = [_issue("unsupported_format", "Document QA поддерживает только PDF и DOCX.")]
        return result

    structural_issues = (
        _docx_structure_issues(target)
        if target.suffix.lower() == ".docx"
        else _pdf_structure_issues(target)
    )

    with tempfile.TemporaryDirectory(prefix="elira-document-qa-") as temp_dir:
        pdf_path, renderer = _render_to_pdf(target, Path(temp_dir))
        result["renderer"] = renderer
        if pdf_path is None:
            result["status"] = "failed" if structural_issues else "unverified"
            result["issues"] = structural_issues + [
                _issue("renderer_unavailable", "Не найден доступный рендерер документа.")
            ]
            return result

        pages = _page_count(pdf_path)
        result["page_count"] = pages
        if pages is None:
            result["status"] = "failed" if structural_issues else "unverified"
            result["issues"] = structural_issues + [
                _issue("page_count_unavailable", "Не удалось определить число страниц PDF после рендера.")
            ]
            return result
        count_issues: list[dict[str, str]] = []
        if expected_page_count is not None and pages != expected_page_count:
            count_issues.append(_issue(
                "page_count_mismatch",
                f"Ожидалось страниц: {expected_page_count}; после рендера: {pages}.",
            ))

        vision_status, vision_issues = _inspect_pages(pdf_path, pages)
        result["vision_status"] = vision_status
        result["issues"] = structural_issues + count_issues + vision_issues
        if structural_issues or count_issues or vision_status == "failed":
            result["status"] = "failed"
        else:
            result["status"] = vision_status
        return result
