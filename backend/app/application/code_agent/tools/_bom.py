"""Deterministic bill-of-materials validation against a local price list."""
from __future__ import annotations

import csv
import hashlib
import json
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Any, Iterator

from app.application.code_agent.tools._sandbox import _resolve_safe

_MONEY = Decimal("0.01")
_MAX_CATALOG_BYTES = 50 * 1024 * 1024
_MAX_CATALOG_ROWS = 100_000
_MAX_CATALOG_COLUMNS = 512
_CATALOG_READ_ERRORS = {
    "sheet_not_found",
    "unsupported_catalog_format",
    "header_row_not_found",
    "empty_header",
    "catalog_too_many_columns",
    "catalog_too_many_rows",
}
_SNAPSHOT_FIELDS = (
    "catalog_sha256",
    "rows",
    "markup_percent",
    "vat_rate",
    "prices_include_vat",
    "subtotal",
    "vat_amount",
    "total",
)


def _money(value: Decimal) -> Decimal:
    return value.quantize(_MONEY, rounding=ROUND_HALF_UP)


def _decimal(value: Any) -> Decimal:
    if isinstance(value, bool) or value is None:
        raise InvalidOperation
    if isinstance(value, (int, float, Decimal)):
        return Decimal(str(value))
    raw = str(value).strip().replace("\u00a0", "").replace(" ", "")
    for token in ("₸", "тг", "KZT", "kzt"):
        raw = raw.replace(token, "")
    if "," in raw and "." not in raw:
        raw = raw.replace(",", ".")
    elif "," in raw and "." in raw:
        raw = raw.replace(",", "")
    return Decimal(raw)


def _positive_quantity(value: Any) -> int:
    number = _decimal(value)
    integral = number.to_integral_value()
    if number != integral or integral <= 0:
        raise InvalidOperation
    return int(integral)


def _bounded_table(
    values: Iterator[tuple[Any, ...] | list[Any]],
    *,
    header_row: int,
    required_columns: set[str],
) -> tuple[list[dict[str, Any]], list[str]]:
    headers: list[str] = []
    indices: dict[str, int] = {}
    rows: list[dict[str, Any]] = []
    scanned_after_header = 0
    for row_number, row in enumerate(values, start=1):
        if row_number < header_row:
            continue
        if row_number == header_row:
            if len(row) > _MAX_CATALOG_COLUMNS:
                raise ValueError("catalog_too_many_columns")
            headers = [str(value or "").strip() for value in row]
            if not any(headers):
                raise ValueError("empty_header")
            indices = {
                header: index
                for index, header in enumerate(headers)
                if header in required_columns
            }
            continue
        scanned_after_header += 1
        if scanned_after_header > _MAX_CATALOG_ROWS:
            raise ValueError("catalog_too_many_rows")
        if not any(value not in (None, "") for value in row):
            continue
        rows.append({
            header: row[index] if index < len(row) else None
            for header, index in indices.items()
        })
    if not headers:
        raise ValueError("header_row_not_found")
    return rows, headers


def _load_rows(
    path: Path,
    sheet_name: str,
    header_row: int,
    required_columns: set[str],
) -> tuple[list[dict[str, Any]], list[str]]:
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            if sheet_name:
                if sheet_name not in workbook.sheetnames:
                    raise ValueError("sheet_not_found")
                sheet = workbook[sheet_name]
            else:
                sheet = workbook.active
            return _bounded_table(
                iter(sheet.iter_rows(values_only=True)),
                header_row=header_row,
                required_columns=required_columns,
            )
        finally:
            workbook.close()
    elif path.suffix.lower() == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            sample = handle.read(8192)
            handle.seek(0)
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
            except csv.Error:
                dialect = csv.excel
            return _bounded_table(
                iter(csv.reader(handle, dialect)),
                header_row=header_row,
                required_columns=required_columns,
            )
    else:
        raise ValueError("unsupported_catalog_format")


def _error(code: str, message: str, **extra: Any) -> dict[str, Any]:
    result = {"ok": False, "error": code, "text": f"ERROR: {message}"}
    result.update(extra)
    return result


def _sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def make_bom_snapshot(result: dict[str, Any]) -> dict[str, Any] | None:
    """Create an immutable, self-verifying snapshot from a successful result."""
    if result.get("ok") is not True or not isinstance(result.get("rows"), list):
        return None
    snapshot = {field: result.get(field) for field in _SNAPSHOT_FIELDS}
    if (
        not snapshot["rows"]
        or not str(snapshot["catalog_sha256"] or "").strip()
        or not str(snapshot["total"] or "").strip()
    ):
        return None
    encoded = json.dumps(
        snapshot,
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
    ).encode("utf-8")
    return {**snapshot, "receipt_sha256": hashlib.sha256(encoded).hexdigest()}


def canonical_bom_file_inputs(
    snapshot: dict[str, Any],
    *,
    format: str,
) -> dict[str, Any]:
    """Build file_gen inputs solely from a verified BOM snapshot."""
    canonical = make_bom_snapshot({"ok": True, **snapshot})
    if (
        canonical is None
        or canonical["receipt_sha256"] != str(snapshot.get("receipt_sha256") or "")
    ):
        raise ValueError("invalid_bom_snapshot")
    headers = ["Код", "Наименование", "Количество", "Цена", "Сумма"]
    data = [
        [
            str(row.get("code") or ""),
            str(row.get("name") or ""),
            int(row.get("quantity") or 0),
            str(row.get("client_unit_price") or ""),
            str(row.get("line_total") or ""),
        ]
        for row in canonical["rows"]
        if isinstance(row, dict)
    ]
    data.extend([
        ["", "Подытог", "", "", str(canonical["subtotal"])],
        ["", f"НДС {canonical['vat_rate']}%", "", "", str(canonical["vat_amount"])],
        ["", "Итого", "", "", str(canonical["total"])],
        ["", "BOM receipt SHA-256", "", "", canonical["receipt_sha256"]],
    ])
    if str(format or "").strip().lower() in {"excel", "xlsx"}:
        return {
            "title": "Спецификация",
            "headers": headers,
            "data": data,
            "content": "",
        }
    lines = ["## Спецификация", " | ".join(headers)]
    lines.extend(" | ".join(str(value) for value in row) for row in data)
    return {
        "title": "Спецификация",
        "headers": None,
        "data": None,
        "content": "\n".join(lines),
    }


def tool_bom_validate(
    project_root: Path,
    *,
    catalog_path: str,
    code_column: str,
    name_column: str,
    price_column: str,
    stock_column: str,
    items: list[dict[str, Any]],
    sheet_name: str = "",
    header_row: int = 1,
    service_items: list[dict[str, Any]] | None = None,
    markup_percent: float = 0,
    vat_rate: float = 12,
    prices_include_vat: bool = True,
    expected_total: float | None = None,
) -> dict[str, Any]:
    """Validate exact catalog codes/stock and calculate canonical client totals."""
    path = _resolve_safe(project_root, str(catalog_path or ""))
    if not path.is_file():
        return _error("catalog_not_found", "catalog file does not exist")
    try:
        if path.stat().st_size > _MAX_CATALOG_BYTES:
            return _error("catalog_too_large", "catalog exceeds the 50 MiB limit")
    except OSError:
        return _error("catalog_read_failed", "could not read the catalog")
    required_columns = {
        str(code_column), str(name_column), str(price_column), str(stock_column),
    }
    try:
        rows, available_columns = _load_rows(
            path,
            str(sheet_name or "").strip(),
            int(header_row),
            required_columns,
        )
    except (OSError, ValueError) as exc:
        code = str(exc) if str(exc) in _CATALOG_READ_ERRORS else "catalog_read_failed"
        return _error(code, "could not read the catalog")
    if not required_columns.issubset(available_columns):
        return _error(
            "catalog_columns_missing",
            "one or more declared catalog columns are missing",
            required_columns=sorted(required_columns),
            available_columns=sorted(available_columns),
        )
    if not isinstance(items, list) or not items:
        return _error("items_empty", "items must contain at least one catalog selection")
    try:
        markup = _decimal(markup_percent)
        vat = _decimal(vat_rate)
    except InvalidOperation:
        return _error("invalid_rate", "markup_percent and vat_rate must be numeric")
    if markup < 0 or vat < 0 or vat > 100:
        return _error("invalid_rate", "markup_percent must be non-negative and vat_rate must be 0..100")

    catalog: dict[str, list[dict[str, Any]]] = {}
    for row in rows:
        code = str(row.get(code_column) or "").strip()
        if code:
            catalog.setdefault(code.casefold(), []).append(row)

    issues: list[dict[str, Any]] = []
    output_rows: list[dict[str, Any]] = []
    selected_codes: set[str] = set()
    subtotal = Decimal("0")
    for selection in items:
        if not isinstance(selection, dict):
            issues.append({"code": "invalid_item", "message": "item must be an object"})
            continue
        code = str(selection.get("code") or "").strip()
        folded = code.casefold()
        if not code:
            issues.append({"code": "item_code_empty", "message": "item code is required"})
            continue
        if folded in selected_codes:
            issues.append({"code": "duplicate_item_code", "item_code": code})
            continue
        selected_codes.add(folded)
        matches = catalog.get(folded, [])
        if not matches:
            issues.append({"code": "catalog_code_not_found", "item_code": code})
            continue
        if len(matches) != 1:
            issues.append({"code": "catalog_code_ambiguous", "item_code": code})
            continue
        row = matches[0]
        try:
            quantity = _positive_quantity(selection.get("quantity"))
        except InvalidOperation:
            issues.append({"code": "invalid_quantity", "item_code": code})
            continue
        try:
            available = _decimal(row.get(stock_column))
            if available != available.to_integral_value() or available < 0:
                raise InvalidOperation
            available_quantity = int(available)
        except InvalidOperation:
            issues.append({"code": "stock_not_numeric", "item_code": code})
            continue
        if quantity > available_quantity:
            issues.append({
                "code": "insufficient_stock",
                "item_code": code,
                "requested_quantity": quantity,
                "available_quantity": available_quantity,
            })
            continue
        try:
            catalog_price = _money(_decimal(row.get(price_column)))
        except InvalidOperation:
            issues.append({"code": "price_not_numeric", "item_code": code})
            continue
        if catalog_price < 0:
            issues.append({"code": "price_negative", "item_code": code})
            continue
        client_price = _money(catalog_price * (Decimal("1") + markup / Decimal("100")))
        line_total = _money(client_price * quantity)
        subtotal += line_total
        output_rows.append({
            "code": code,
            "name": str(row.get(name_column) or "").strip(),
            "quantity": quantity,
            "available_quantity": available_quantity,
            "catalog_unit_price": f"{catalog_price:.2f}",
            "client_unit_price": f"{client_price:.2f}",
            "line_total": f"{line_total:.2f}",
            "source": "catalog",
        })

    for service in service_items or []:
        if not isinstance(service, dict):
            issues.append({"code": "invalid_service", "message": "service item must be an object"})
            continue
        code = str(service.get("code") or "").strip()
        name = str(service.get("name") or "").strip()
        try:
            quantity = _positive_quantity(service.get("quantity"))
            unit_price = _money(_decimal(service.get("unit_price")))
        except InvalidOperation:
            issues.append({"code": "invalid_service_amount", "item_code": code})
            continue
        if not code or not name or unit_price < 0:
            issues.append({"code": "invalid_service", "item_code": code})
            continue
        line_total = _money(unit_price * quantity)
        subtotal += line_total
        output_rows.append({
            "code": code,
            "name": name,
            "quantity": quantity,
            "client_unit_price": f"{unit_price:.2f}",
            "line_total": f"{line_total:.2f}",
            "source": "service",
        })

    if issues:
        result = {
            "ok": False,
            "error": "bom_validation_failed",
            "issues": issues,
            "validated_rows": output_rows,
        }
        result["text"] = json.dumps(result, ensure_ascii=False, separators=(",", ":"))
        return result

    subtotal = _money(subtotal)
    if prices_include_vat:
        vat_amount = _money(subtotal * vat / (Decimal("100") + vat)) if vat else Decimal("0.00")
        total = subtotal
    else:
        vat_amount = _money(subtotal * vat / Decimal("100"))
        total = _money(subtotal + vat_amount)
    result: dict[str, Any] = {
        "ok": True,
        "catalog_sha256": _sha256_file(path),
        "rows": output_rows,
        "markup_percent": f"{markup:.2f}",
        "vat_rate": f"{vat:.2f}",
        "prices_include_vat": bool(prices_include_vat),
        "subtotal": f"{subtotal:.2f}",
        "vat_amount": f"{vat_amount:.2f}",
        "total": f"{total:.2f}",
    }
    if expected_total is not None:
        try:
            expected = _money(_decimal(expected_total))
        except InvalidOperation:
            return _error("expected_total_invalid", "expected_total must be numeric")
        result["expected_total"] = f"{expected:.2f}"
        result["expected_total_matches"] = expected == total
        if expected != total:
            result["ok"] = False
            result["error"] = "expected_total_mismatch"
    snapshot = make_bom_snapshot(result)
    if snapshot is not None:
        result["receipt_sha256"] = snapshot["receipt_sha256"]
    result["text"] = json.dumps(result, ensure_ascii=False, separators=(",", ":"))
    return result
