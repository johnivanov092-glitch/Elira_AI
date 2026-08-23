"""
skills_service.py — 4 скилла Elira.

1. Генерация Word/Excel
2. SQL запросы (SQLite)
3. HTTP/API вызовы
4. Скриншот сайта (playwright)
"""
from __future__ import annotations
import html
import logging
import re
import sqlite3
import time
from pathlib import Path
from typing import Any

import requests as http_lib

from app.core.config import DATA_DIR, GENERATED_DIR
from app.infrastructure.db.connection import connect_sqlite

logger = logging.getLogger(__name__)

OUTPUT_DIR = GENERATED_DIR
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)


def screenshot_capability_status() -> dict:
    try:
        from playwright.sync_api import sync_playwright  # noqa: F401
    except ImportError:
        return {
            "feature": "screenshot",
            "available": False,
            "reason": "optional_dependency_missing",
            "missing_packages": ["playwright"],
            "hint": "pip install playwright && playwright install chromium",
        }
    return {
        "feature": "screenshot",
        "available": True,
        "reason": None,
        "missing_packages": [],
        "hint": None,
    }


# ═══════════════════════════════════════════════════════════════
# 1. ГЕНЕРАЦИЯ ФАЙЛОВ
# ═══════════════════════════════════════════════════════════════

def generate_word(title: str, content: str, filename: str = "") -> dict:
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
    except ImportError:
        return {"ok": False, "error": "pip install python-docx"}

    doc = Document()
    if title:
        h = doc.add_heading(title, level=1)
        h.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for line in content.split("\n"):
        line = line.strip()
        if not line:
            doc.add_paragraph("")
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=2)
        elif line.startswith("### "):
            doc.add_heading(line[4:], level=3)
        elif line.startswith("- ") or line.startswith("* "):
            doc.add_paragraph(line[2:], style="List Bullet")
        elif line[0:3] in ("1. ", "2. ", "3. ", "4. ", "5. ", "6. ", "7. ", "8. ", "9. "):
            doc.add_paragraph(line[3:], style="List Number")
        else:
            doc.add_paragraph(line)

    fname = filename or f"elira_{int(time.time())}.docx"
    if not fname.endswith(".docx"):
        fname += ".docx"
    path = OUTPUT_DIR / fname
    doc.save(str(path))
    return {"ok": True, "path": str(path), "filename": fname, "size": path.stat().st_size,
            "download_url": f"/api/skills/download/{fname}"}


def generate_excel(title: str, data: list, headers: list = None, filename: str = "") -> dict:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        return {"ok": False, "error": "pip install openpyxl"}

    wb = Workbook()
    ws = wb.active
    ws.title = title or "Sheet1"

    if headers:
        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.font = Font(bold=True, size=11)
            cell.fill = PatternFill(start_color="D5E8F0", end_color="D5E8F0", fill_type="solid")
            cell.alignment = Alignment(horizontal="center")

    start_row = 2 if headers else 1
    for row_idx, row_data in enumerate(data, start_row):
        if isinstance(row_data, (list, tuple)):
            for col_idx, value in enumerate(row_data, 1):
                ws.cell(row=row_idx, column=col_idx, value=value)

    for col in ws.columns:
        max_len = max((len(str(cell.value or "")) for cell in col), default=8)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 50)

    fname = filename or f"elira_{int(time.time())}.xlsx"
    if not fname.endswith(".xlsx"):
        fname += ".xlsx"
    path = OUTPUT_DIR / fname
    wb.save(str(path))
    return {"ok": True, "path": str(path), "filename": fname, "size": path.stat().st_size,
            "download_url": f"/api/skills/download/{fname}"}


# ─── PDF (fixed static template → headless Chromium print) ───────────────────
# Production-safe: a FIXED A4 HTML/CSS shell. title/content are html.escape'd and
# rendered as PLAIN TEXT (white-space: pre-wrap keeps line breaks). NO raw HTML from
# the model, NO URLs / external resources / <script> — so nothing the content carries
# can execute or fetch. Reuses the already-installed Playwright + Chromium (no new
# dependency, no shell, no external converter).
_PDF_HTML_TEMPLATE = (
    "<!doctype html><html><head><meta charset=\"utf-8\"><style>"
    "@page {{ size: A4; margin: 20mm; }}"
    "html, body {{ font-family: 'Segoe UI', Arial, sans-serif; color: #111; }}"
    "h1 {{ font-size: 20pt; margin: 0 0 12pt 0; }}"
    ".content {{ font-size: 12pt; line-height: 1.5; white-space: pre-wrap; word-wrap: break-word; }}"
    "</style></head><body>{title_block}<div class=\"content\">{content}</div></body></html>"
)


def _safe_pdf_basename(filename: str) -> str:
    """Safe .pdf filename from caller input: basename ONLY (no path separators /
    traversal), server-owned extension. NOTE: only the NEW PDF branch is hardened
    here — the legacy Word/Excel filename paths are intentionally left as-is (that
    sanitization debt is tracked separately)."""
    raw = (filename or "").strip()
    base = re.split(r"[\\/]", raw)[-1].strip()          # drop any directory part
    if base.lower().endswith(".pdf"):
        base = base[:-4]
    base = base.strip().strip(".")                       # no leading/trailing dots ('..')
    base = re.sub(r"[^\w.\- ]", "_", base).strip()       # conservative allowlist (keeps unicode \w)
    if not base:
        base = f"elira_{int(time.time())}"
    return f"{base}.pdf"


def generate_pdf(title: str, content: str, filename: str = "") -> dict:
    try:
        from playwright.sync_api import sync_playwright
    except ImportError:
        return {"ok": False, "error": "pip install playwright && playwright install chromium"}

    fname = _safe_pdf_basename(filename)
    path = OUTPUT_DIR / fname

    title_block = f"<h1>{html.escape(title)}</h1>" if title else ""
    doc_html = _PDF_HTML_TEMPLATE.format(
        title_block=title_block, content=html.escape(content or "")
    )

    try:
        with sync_playwright() as p:
            browser = p.chromium.launch(headless=True)
            try:
                page = browser.new_page()
                page.set_content(doc_html, wait_until="load", timeout=15000)
                page.pdf(path=str(path), format="A4", print_background=True)
            finally:
                browser.close()
    except Exception as e:
        return {"ok": False, "error": str(e)}

    if not path.exists() or path.stat().st_size == 0:
        return {"ok": False, "error": "PDF generation produced no output file"}
    return {"ok": True, "path": str(path), "filename": fname, "size": path.stat().st_size,
            "download_url": f"/api/skills/download/{fname}"}


# ═══════════════════════════════════════════════════════════════
# 2. SQL ЗАПРОСЫ
# ═══════════════════════════════════════════════════════════════

def _safe_db(db_path: str) -> Path:
    """Resolve a database path without imposing a product filesystem scope."""
    raw_text = str(db_path or "").strip()
    if not raw_text:
        raise ValueError("Путь к базе данных не указан")
    raw = Path(raw_text).expanduser()
    return (raw if raw.is_absolute() else DATA_DIR / raw).resolve()


def run_sql(db_path: str, query: str, params: list = None, max_rows: int = 100) -> dict:
    try:
        safe = _safe_db(db_path)
    except ValueError as e:
        return {"ok": False, "error": str(e)}
    if not safe.exists():
        return {"ok": False, "error": f"Не найдена: {db_path}"}

    q_up = query.strip().upper()

    try:
        conn = connect_sqlite(safe, row_factory=sqlite3.Row, journal_mode=None)
        cur = conn.cursor()
        if q_up.startswith(("SELECT", "PRAGMA", "WITH")):
            cur.execute(query, params or [])
            rows = cur.fetchmany(max_rows)
            columns = [d[0] for d in cur.description] if cur.description else []
            data = [dict(r) for r in rows]
            conn.close()
            return {"ok": True, "columns": columns, "rows": data, "count": len(data)}
        else:
            cur.execute(query, params or [])
            conn.commit()
            affected = cur.rowcount
            conn.close()
            return {"ok": True, "affected": affected}
    except Exception as e:
        return {"ok": False, "error": str(e)}


def list_databases() -> dict:
    dbs = []
    for d in [DATA_DIR.resolve()]:
        if d.exists():
            for f in d.rglob("*.db"):
                dbs.append({"path": str(f), "name": f.name, "size": f.stat().st_size})
    return {"ok": True, "databases": dbs}


def describe_db(db_path: str) -> dict:
    try:
        safe = _safe_db(db_path)
    except ValueError as e:
        return {"ok": False, "error": str(e)}
    try:
        conn = connect_sqlite(safe, row_factory=sqlite3.Row, journal_mode=None)
        tables = conn.execute("SELECT name FROM sqlite_master WHERE type='table' ORDER BY name").fetchall()
        schema = {}
        for (tbl,) in tables:
            safe_tbl = '"' + tbl.replace('"', '""') + '"'
            cols = conn.execute(f"PRAGMA table_info({safe_tbl})").fetchall()
            cnt = conn.execute(f"SELECT COUNT(*) FROM {safe_tbl}").fetchone()[0]
            schema[tbl] = {"columns": [{"name": c[1], "type": c[2], "pk": bool(c[5])} for c in cols], "rows": cnt}
        conn.close()
        return {"ok": True, "tables": schema}
    except Exception as e:
        return {"ok": False, "error": str(e)}


# ═══════════════════════════════════════════════════════════════
# 3. HTTP / API
# ═══════════════════════════════════════════════════════════════

def http_request(url: str, method: str = "GET", headers: dict = None, body: Any = None,
                 timeout: int = 15, allow_loopback_ports: set = None) -> dict:
    del allow_loopback_ports

    try:
        kw = {"url": url, "headers": headers or {}, "timeout": timeout}
        method = method.upper()
        if method == "GET":
            resp = http_lib.get(**kw)
        elif method == "POST":
            kw["json"] = body if isinstance(body, (dict, list)) else None
            kw["data"] = body if not isinstance(body, (dict, list)) else None
            resp = http_lib.post(**kw)
        elif method == "PUT":
            kw["json"] = body if isinstance(body, (dict, list)) else None
            resp = http_lib.put(**kw)
        elif method == "DELETE":
            resp = http_lib.delete(**kw)
        else:
            return {"ok": False, "error": f"Неизвестный метод: {method}"}

        ct = resp.headers.get("content-type", "")
        try:
            rbody = resp.json() if "json" in ct else resp.text[:30000]
        except Exception:
            rbody = resp.text[:30000]

        return {"ok": True, "status": resp.status_code, "body": rbody,
                "url": str(resp.url), "elapsed_ms": int(resp.elapsed.total_seconds() * 1000)}
    except http_lib.Timeout:
        return {"ok": False, "error": f"Таймаут ({timeout}с)"}
    except Exception as e:
        return {"ok": False, "error": str(e)}


# ═══════════════════════════════════════════════════════════════
# 4. СКРИНШОТ САЙТА
# ═══════════════════════════════════════════════════════════════

def screenshot_url(url: str, width: int = 1280, height: int = 800, full_page: bool = False) -> dict:
    try:
        from playwright.sync_api import sync_playwright
    except ImportError:
        status = screenshot_capability_status()
        return {
            "ok": False,
            "error": "Screenshot feature is unavailable",
            "feature": status["feature"],
            "reason": status["reason"],
            "missing_packages": status["missing_packages"],
            "hint": status["hint"],
        }

    fname = f"screenshot_{int(time.time())}.png"
    path = OUTPUT_DIR / fname
    try:
        with sync_playwright() as p:
            browser = p.chromium.launch(headless=True)
            page = browser.new_page(viewport={"width": width, "height": height})
            page.goto(url, wait_until="networkidle", timeout=20000)
            page.screenshot(path=str(path), full_page=full_page)
            title = page.title()
            browser.close()
        return {"ok": True, "path": str(path), "filename": fname, "title": title,
                "download_url": f"/api/skills/download/{fname}", "view_url": f"/api/skills/view/{fname}"}
    except Exception as e:
        return {"ok": False, "error": str(e)}
