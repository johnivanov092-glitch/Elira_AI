"""Create DOCX/XLSX/PDF from a JSON file without application code or publication."""
from __future__ import annotations

import argparse
import hashlib
import html
import json
import os
from pathlib import Path
import re
import sys
import tempfile
from typing import Any


def require_text(payload: dict[str, Any], name: str) -> str:
    value = payload.get(name, "")
    if not isinstance(value, str):
        raise ValueError(f"{name} must be a string")
    return value


def rows(value: Any, name: str) -> list[list[Any]]:
    if not isinstance(value, list) or any(not isinstance(row, list) for row in value):
        raise ValueError(f"{name} must be a list of rows")
    if any(not isinstance(cell, (str, int, float, bool, type(None))) for row in value for cell in row):
        raise ValueError(f"{name} cells must be strings, numbers, booleans or null")
    return value


def create_docx(payload: dict[str, Any], path: Path) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    document = Document()
    title = require_text(payload, "title")
    if title:
        document.add_heading(title, 1).alignment = WD_ALIGN_PARAGRAPH.CENTER
    for line in require_text(payload, "content").splitlines():
        if line.startswith("### "):
            document.add_heading(line[4:], 3)
        elif line.startswith("## "):
            document.add_heading(line[3:], 2)
        elif line.startswith(("- ", "* ")):
            document.add_paragraph(line[2:], style="List Bullet")
        elif re.match(r"^\d+\. ", line):
            document.add_paragraph(line.split(". ", 1)[1], style="List Number")
        else:
            document.add_paragraph(line)
    tables = payload.get("tables", [])
    if not isinstance(tables, list):
        raise ValueError("tables must be a list of tables")
    for values in tables:
        matrix = rows(values, "tables")
        width = max((len(row) for row in matrix), default=0)
        if not width:
            continue
        table = document.add_table(rows=0, cols=width)
        table.style = "Table Grid"
        for row in matrix:
            cells = table.add_row().cells
            for column, value in enumerate(row):
                cells[column].text = "" if value is None else str(value)
    document.save(path)


def create_xlsx(payload: dict[str, Any], path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = require_text(payload, "title") or "Sheet1"
    headers = payload.get("headers", [])
    if not isinstance(headers, list) or any(not isinstance(value, str) for value in headers):
        raise ValueError("headers must be a list of strings")
    formula_flag = payload.get("allow_formulas", False)
    if not isinstance(formula_flag, bool):
        raise ValueError("allow_formulas must be a boolean")
    values = rows(payload.get("data", []), "data")
    if headers:
        values = [headers, *values]
    for row_number, row in enumerate(values, 1):
        for column, value in enumerate(row, 1):
            cell = sheet.cell(row_number, column, value)
            if isinstance(value, str) and not formula_flag:
                cell.data_type = "s"
            if row_number == 1 and headers:
                cell.font = Font(bold=True)
                cell.fill = PatternFill(fill_type="solid", fgColor="D5E8F0")
                cell.alignment = Alignment(horizontal="center")
    for column in sheet.columns:
        width = max((len(str(cell.value or "")) for cell in column), default=8)
        sheet.column_dimensions[column[0].column_letter].width = min(width + 4, 50)
    workbook.save(path)
    workbook.close()


def unicode_font(explicit: Path | None) -> Path:
    candidates = [explicit] if explicit else [
        Path("C:/Windows/Fonts/arial.ttf"), Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
        Path("/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf"),
    ]
    for candidate in candidates:
        if candidate and candidate.is_file():
            return candidate
    raise ValueError("Unicode TTF font not found; pass --font PATH")


def create_pdf(payload: dict[str, Any], path: Path, font: Path | None) -> None:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
    pdfmetrics.registerFont(TTFont("SkillUnicode", str(unicode_font(font))))
    normal = ParagraphStyle("Body", fontName="SkillUnicode", fontSize=11, leading=16, textColor=colors.black)
    heading = ParagraphStyle("Title", parent=normal, fontSize=18, leading=24, alignment=TA_CENTER, spaceAfter=12)
    story = []
    title = require_text(payload, "title")
    if title:
        story.append(Paragraph(html.escape(title), heading))
    for line in require_text(payload, "content").splitlines():
        story.append(Paragraph(html.escape(line), normal) if line else Spacer(1, 8))
    if not story:
        story.append(Spacer(1, 8))
    SimpleDocTemplate(str(path), pagesize=A4, rightMargin=20*mm, leftMargin=20*mm,
                      topMargin=20*mm, bottomMargin=20*mm).build(story)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--format", required=True, choices=("docx", "xlsx", "pdf"))
    parser.add_argument("--input", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--font", type=Path)
    parser.add_argument("--overwrite", action="store_true")
    args = parser.parse_args()
    temporary = None
    try:
        if args.output.suffix.lower() != "." + args.format:
            raise ValueError("Output extension must match --format")
        if args.output.resolve() == args.input.resolve():
            raise ValueError("Output must not overwrite the JSON input")
        if args.output.exists() and not args.overwrite:
            raise ValueError("Output exists; choose another path or explicitly use --overwrite")
        payload = json.loads(args.input.read_text(encoding="utf-8-sig"))
        if not isinstance(payload, dict):
            raise ValueError("Input JSON must be an object")
        with tempfile.NamedTemporaryFile(dir=args.output.parent, suffix=args.output.suffix, delete=False) as handle:
            temporary = Path(handle.name)
        if args.format == "docx":
            create_docx(payload, temporary)
        elif args.format == "xlsx":
            create_xlsx(payload, temporary)
        else:
            create_pdf(payload, temporary, args.font)
        if not temporary.stat().st_size:
            raise ValueError("Document generator produced an empty file")
        if args.overwrite:
            os.replace(temporary, args.output)
        else:
            # Exclusive create also prevents a racing writer from being overwritten.
            with args.output.open("xb") as destination, temporary.open("rb") as source:
                import shutil
                shutil.copyfileobj(source, destination)
        print(json.dumps({"ok": True, "path": str(args.output.resolve()), "format": args.format,
                          "size": args.output.stat().st_size,
                          "sha256": hashlib.sha256(args.output.read_bytes()).hexdigest()}, ensure_ascii=False))
        return 0
    except Exception as exc:
        print(json.dumps({"ok": False, "error": type(exc).__name__, "message": str(exc)}, ensure_ascii=False))
        return 1
    finally:
        if temporary:
            temporary.unlink(missing_ok=True)


if __name__ == "__main__":
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    raise SystemExit(main())
