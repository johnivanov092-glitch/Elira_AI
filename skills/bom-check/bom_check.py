"""Проверка спецификации (BOM/КП) по локальному прайсу XLSX/CSV.

Запуск: python bom_check.py spec.json
spec.json:
{
  "catalog_path": "прайс.xlsx",          # XLSX/XLSM или CSV
  "sheet_name": "",                      # лист XLSX, по умолчанию активный
  "header_row": 1,                       # строка заголовков (с 1)
  "code_column": "Код", "name_column": "Наименование",
  "price_column": "Цена", "stock_column": "Остаток",
  "items": [{"code": "A-100", "quantity": 2}],
  "service_items": [{"code": "S1", "name": "Сборка", "quantity": 1, "unit_price": 5000}],
  "markup_percent": 0,
  "vat_rate": 16,                        # НДС РК с 01.01.2026
  "prices_include_vat": true,
  "expected_total": null
}
Печатает таблицу и итоги; последней строкой — JSON с теми же числами.
Код выхода: 0 — всё сошлось, 1 — есть проблемы (коды не найдены, остатка мало,
итог не совпал), 2 — ошибка входных данных.
"""
from __future__ import annotations

import csv
import hashlib
import json
import sys
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Iterator

MAX_ROWS = 100_000
_CURRENCY = ("₸", "тг.", "тг", "тенге", "KZT", "kzt", "₽", "руб.", "руб", "RUB", "$", "USD", "€", "EUR")


def _plain_number(raw: str) -> str:
    """Separators of a price-list number -> a plain decimal string.

    The last separator of two different ones is the decimal point ("1.234,56",
    "1,234.56"); a separator repeated twice or more groups thousands ("1 234 567"
    after spaces, "1.234.567"). A single separator followed by exactly three
    digits after a non-zero integer part ("1,234", "12.500") may be either and
    is refused — a wrong guess would scale the price by 1000.
    """
    commas, dots = raw.count(","), raw.count(".")
    if commas and dots:
        decimal_sep = "," if raw.rfind(",") > raw.rfind(".") else "."
        group_sep = "." if decimal_sep == "," else ","
        if raw.count(decimal_sep) > 1:
            raise InvalidOperation
        return raw.replace(group_sep, "").replace(decimal_sep, ".")
    sep = "," if commas else "." if dots else ""
    if not sep:
        return raw
    if raw.count(sep) > 1:
        return raw.replace(sep, "")
    whole, _, fraction = raw.partition(sep)
    if len(fraction) == 3 and whole.lstrip("+-").lstrip("0"):
        raise InvalidOperation
    return whole + "." + fraction


def number(value: Any) -> Decimal:
    """Число как в прайсе: пробелы, валюта, запятая как дробная часть."""
    if isinstance(value, bool) or value is None:
        raise InvalidOperation
    if isinstance(value, (int, float, Decimal)):
        result = Decimal(str(value))
    else:
        raw = str(value).strip().replace(" ", "").replace(" ", "").replace(" ", "")
        for token in _CURRENCY:
            raw = raw.replace(token, "")
        raw = _plain_number(raw)
        result = Decimal(raw)
    if not result.is_finite():
        raise InvalidOperation
    return result


def money(value: Decimal) -> Decimal:
    return value.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def quantity(value: Any) -> int:
    n = number(value)
    if n != n.to_integral_value() or n <= 0:
        raise InvalidOperation
    return int(n)


def table(rows: Iterator, header_row: int, columns: set[str]) -> tuple[list[dict], list[str]]:
    headers: list[str] = []
    index: dict[str, int] = {}
    out: list[dict] = []
    for n, row in enumerate(rows, start=1):
        if n < header_row:
            continue
        if n == header_row:
            headers = [str(v or "").strip() for v in row]
            index = {h: i for i, h in enumerate(headers) if h in columns}
            continue
        if len(out) >= MAX_ROWS:
            raise ValueError(f"в прайсе больше {MAX_ROWS} строк")
        if any(v not in (None, "") for v in row):
            out.append({h: row[i] if i < len(row) else None for h, i in index.items()})
    if not headers:
        raise ValueError(f"строка заголовков {header_row} не найдена")
    return out, headers


def load(path: Path, sheet: str, header_row: int, columns: set[str]) -> tuple[list[dict], list[str]]:
    suffix = path.suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook

        book = load_workbook(path, read_only=True, data_only=True)
        try:
            if sheet and sheet not in book.sheetnames:
                raise ValueError(f"лист «{sheet}» не найден; есть: {', '.join(book.sheetnames)}")
            ws = book[sheet] if sheet else book.active
            return table(iter(ws.iter_rows(values_only=True)), header_row, columns)
        finally:
            book.close()
    if suffix == ".csv":
        raw = path.read_bytes()
        for encoding in ("utf-8-sig", "cp1251"):
            try:
                text = raw.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        else:
            raise ValueError("CSV не в UTF-8 и не в cp1251")
        try:
            dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        return table(iter(csv.reader(text.splitlines(), dialect)), header_row, columns)
    raise ValueError("прайс должен быть .xlsx, .xlsm или .csv")


def check(spec: dict[str, Any], base: Path) -> dict[str, Any]:
    path = Path(str(spec.get("catalog_path") or ""))
    if not path.is_absolute():
        path = base / path
    if not path.is_file():
        raise ValueError(f"прайс не найден: {path}")
    cols = {k: str(spec.get(k) or "") for k in ("code_column", "name_column", "price_column", "stock_column")}
    if not all(cols.values()):
        raise ValueError("укажи code_column, name_column, price_column и stock_column")
    rows, headers = load(path, str(spec.get("sheet_name") or ""), int(spec.get("header_row") or 1),
                         set(cols.values()))
    missing = sorted(set(cols.values()) - set(headers))
    if missing:
        raise ValueError(f"в прайсе нет колонок {missing}; есть: {headers}")
    items = spec.get("items") or []
    if not isinstance(items, list) or not items:
        raise ValueError("items: нужен хотя бы один код")
    markup = number(spec.get("markup_percent") or 0)
    vat = number(spec.get("vat_rate") if spec.get("vat_rate") is not None else 16)
    if markup < 0 or not (0 <= vat <= 100):
        raise ValueError("markup_percent ≥ 0, vat_rate 0..100")

    catalog: dict[str, list[dict]] = {}
    for row in rows:
        code = str(row.get(cols["code_column"]) or "").strip()
        if code:
            catalog.setdefault(code.casefold(), []).append(row)

    issues: list[str] = []
    lines: list[dict[str, Any]] = []
    seen: set[str] = set()
    subtotal = Decimal("0")
    for item in items:
        code = str((item or {}).get("code") or "").strip()
        if not code or code.casefold() in seen:
            issues.append(f"{code or '(пусто)'}: пустой или повторный код")
            continue
        seen.add(code.casefold())
        matches = catalog.get(code.casefold(), [])
        if len(matches) != 1:
            issues.append(f"{code}: {'нет в прайсе' if not matches else 'несколько строк в прайсе'}")
            continue
        row = matches[0]
        try:
            qty = quantity(item.get("quantity"))
            stock = number(row.get(cols["stock_column"]))
            price = money(number(row.get(cols["price_column"])))
        except InvalidOperation:
            issues.append(f"{code}: количество, остаток или цена не число или записаны неоднозначно (1,234 — тысяча или дробь?)")
            continue
        if qty > stock:
            issues.append(f"{code}: нужно {qty}, на остатке {stock}")
            continue
        unit = money(price * (1 + markup / 100))
        total = money(unit * qty)
        subtotal += total
        lines.append({"code": code, "name": str(row.get(cols["name_column"]) or "").strip(),
                      "quantity": qty, "catalog_price": f"{price:.2f}",
                      "unit_price": f"{unit:.2f}", "line_total": f"{total:.2f}"})
    for service in spec.get("service_items") or []:
        code = str((service or {}).get("code") or "").strip()
        name = str((service or {}).get("name") or "").strip()
        try:
            qty = quantity(service.get("quantity"))
            unit = money(number(service.get("unit_price")))
        except InvalidOperation:
            issues.append(f"{code or name}: услуга — количество или цена не число")
            continue
        total = money(unit * qty)
        subtotal += total
        lines.append({"code": code, "name": name, "quantity": qty, "catalog_price": "",
                      "unit_price": f"{unit:.2f}", "line_total": f"{total:.2f}", "service": True})

    subtotal = money(subtotal)
    if spec.get("prices_include_vat", True):
        vat_amount = money(subtotal * vat / (100 + vat)) if vat else Decimal("0.00")
        grand = subtotal
    else:
        vat_amount = money(subtotal * vat / 100)
        grand = money(subtotal + vat_amount)
    result: dict[str, Any] = {
        "ok": not issues, "issues": issues, "rows": lines,
        "markup_percent": f"{markup:.2f}", "vat_rate": f"{vat:.2f}",
        "prices_include_vat": bool(spec.get("prices_include_vat", True)),
        "subtotal": f"{subtotal:.2f}", "vat_amount": f"{vat_amount:.2f}", "total": f"{grand:.2f}",
        "catalog_sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
    }
    expected = spec.get("expected_total")
    if expected is not None:
        if money(number(expected)) != grand:
            result["ok"] = False
            issues.append(f"итог {grand:.2f} не совпал с ожидаемым {money(number(expected)):.2f}")
    return result


def main() -> int:
    for stream in (sys.stdout, sys.stderr):
        if hasattr(stream, "reconfigure"):
            stream.reconfigure(encoding="utf-8")
    if len(sys.argv) != 2:
        print(__doc__)
        return 2
    spec_path = Path(sys.argv[1])
    try:
        spec = json.loads(spec_path.read_text(encoding="utf-8-sig"))
        result = check(spec, spec_path.resolve().parent)
    except (OSError, ValueError, InvalidOperation, json.JSONDecodeError) as exc:
        print(f"ОШИБКА: {exc}")
        return 2
    for row in result["rows"]:
        print(f"{row['code']:<14} {row['name'][:40]:<40} {row['quantity']:>5} x {row['unit_price']:>12} = {row['line_total']:>12}")
    vat_word = "в т.ч. НДС" if result["prices_include_vat"] else "НДС сверху"
    print(f"Подытог {result['subtotal']}; {vat_word} {result['vat_rate']}% = {result['vat_amount']}; Итого {result['total']}")
    for issue in result["issues"]:
        print(f"ПРОБЛЕМА: {issue}")
    print(json.dumps(result, ensure_ascii=False))
    return 0 if result["ok"] else 1


if __name__ == "__main__":
    sys.exit(main())
