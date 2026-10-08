"""Standalone document reader. No application imports, OCR or automatic uploads."""
from __future__ import annotations

import argparse
import hashlib
import json
import sys
import tempfile
import time
import zipfile
from pathlib import Path
from typing import Any
from urllib.parse import urljoin, urlsplit


TEXT_EXTENSIONS = {
    ".txt", ".md", ".csv", ".tsv", ".json", ".xml", ".html", ".htm",
    ".log", ".yaml", ".yml", ".toml", ".ini", ".cfg", ".conf",
    ".py", ".js", ".ts", ".css", ".sql", ".bas", ".vba", ".vbs",
    ".cls", ".frm", ".rsc", ".ps1", ".sh", ".bat", ".cmd",
    ".jsx", ".tsx", ".env", ".rb", ".php", ".java", ".c", ".cpp", ".h", ".hpp",
    ".cs", ".go", ".rs", ".swift", ".kt", ".r", ".m", ".lua", ".pl", ".tcl", ".asm",
    ".gitignore", ".dockerfile", ".makefile", ".sln", ".csproj", ".pom", ".gradle",
}
MAX_WORKSHEET_CELLS = 2_000_000


def decode_text(data: bytes, encoding: str = "utf-8-sig") -> str:
    try:
        return data.decode(encoding)
    except UnicodeDecodeError as exc:
        raise ValueError("Text is not valid in the selected encoding; explicitly use --encoding cp1251/cp866 if appropriate") from exc


def download_document(url: str, destination: Path, timeout: float, max_bytes: int) -> dict[str, str]:
    """Explicit HTTP download, finite redirects/size, retained original and provenance."""
    import requests
    if destination.exists():
        raise ValueError("Download destination exists; choose another path")
    requested, current = url, url
    deadline = time.monotonic() + timeout
    temporary = None
    try:
        for redirect in range(6):
            if time.monotonic() >= deadline:
                raise requests.Timeout("Document download deadline exceeded")
            parsed = urlsplit(current)
            if parsed.scheme not in {"http", "https"} or not parsed.hostname or parsed.username or parsed.password:
                raise ValueError("Document URL must use HTTP(S), without embedded credentials")
            with requests.get(current, stream=True, allow_redirects=False, timeout=(min(10, timeout), timeout)) as response:
                if 300 <= response.status_code < 400:
                    location = response.headers.get("Location")
                    if not location or redirect == 5:
                        raise ValueError("Document redirect is missing or exceeds five redirects")
                    current = urljoin(current, location)
                    continue
                response.raise_for_status()
                declared = response.headers.get("Content-Length")
                if declared is not None and int(declared) > max_bytes:
                    raise ValueError("Document exceeds max-bytes")
                digest, total = hashlib.sha256(), 0
                with tempfile.NamedTemporaryFile(dir=destination.parent, delete=False) as handle:
                    temporary = Path(handle.name)
                    for chunk in response.iter_content(chunk_size=65536):
                        if time.monotonic() >= deadline:
                            raise requests.Timeout("Document download deadline exceeded")
                        total += len(chunk)
                        if total > max_bytes:
                            raise ValueError("Document exceeds max-bytes")
                        digest.update(chunk)
                        handle.write(chunk)
                if total == 0:
                    raise ValueError("Downloaded document is empty")
                with destination.open("xb") as output, temporary.open("rb") as source:
                    import shutil
                    shutil.copyfileobj(source, output)
                return {"requested_url": requested, "final_url": str(response.url),
                        "sha256": digest.hexdigest(), "local_path": str(destination.resolve())}
        raise ValueError("Document redirect limit exceeded")
    finally:
        if temporary:
            temporary.unlink(missing_ok=True)


def read_pdf(path: Path) -> dict[str, Any]:
    from pypdf import PdfReader
    import pdfplumber

    reader = PdfReader(path)
    if reader.is_encrypted and not reader.decrypt(""):
        raise ValueError("PDF is encrypted; supply an unencrypted copy")
    pages, tables, errors = [], [], []
    with pdfplumber.open(path) as plumber:
        for number, page in enumerate(reader.pages, 1):
            text, method = "", "pypdf"
            try:
                text = page.extract_text() or ""
            except Exception as exc:
                errors.append({"page": number, "code": "native_text_failed", "message": str(exc)})
            source = plumber.pages[number - 1]
            if not text.strip():
                try:
                    text = source.extract_text() or ""
                    method = "pdfplumber"
                except Exception as exc:
                    errors.append({"page": number, "code": "fallback_text_failed", "message": str(exc)})
            try:
                for index, rows in enumerate(source.extract_tables() or [], 1):
                    tables.append({"page": number, "index": index, "rows": rows})
            except Exception as exc:
                errors.append({"page": number, "code": "table_extraction_failed", "message": str(exc)})
            pages.append({"page": number, "text": text, "method": method, "chars": len(text)})
    missing = [item["page"] for item in pages if not item["text"].strip()]
    return {"text": "\n\n".join(f"=== Страница {p['page']} ===\n{p['text']}" for p in pages),
            "pages": pages, "tables": tables, "ocr_required": missing, "errors": errors,
            "complete": not missing and not errors}


def read_document(path: Path, encoding: str = "utf-8-sig") -> dict[str, Any]:
    extension = path.suffix.lower()
    result: dict[str, Any] = {"text": "", "tables": [], "errors": [], "ocr_required": [], "complete": True}
    if extension == ".pdf":
        result.update(read_pdf(path))
    elif extension == ".docx":
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph
        document = Document(path)
        parts = []
        for block in document.iter_inner_content():
            if isinstance(block, Paragraph):
                parts.append(block.text)
            elif isinstance(block, Table):
                index = len(result["tables"]) + 1
                rows = [[cell.text for cell in row.cells] for row in block.rows]
                result["tables"].append({"index": index, "rows": rows})
                parts.append(f"=== Таблица {index} ===\n" + "\n".join(" | ".join(row) for row in rows))
        for index, section in enumerate(document.sections, 1):
            for label, container in (("Колонтитул сверху", section.header), ("Колонтитул снизу", section.footer)):
                value = "\n".join(p.text for p in container.paragraphs if p.text.strip())
                if value:
                    parts.append(f"=== {label} {index} ===\n{value}")
        result["text"] = "\n".join(parts)
    elif extension in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=True)
        formulas = load_workbook(path, read_only=True, data_only=False)
        try:
            for sheet in workbook.worksheets:
                if (sheet.max_row or 0) * (sheet.max_column or 0) > MAX_WORKSHEET_CELLS:
                    raise ValueError(f"Worksheet {sheet.title!r} declares too many cells; explicitly reduce its dimensions before reading")
            result["sheets"] = [{"name": sheet.title, "rows": [list(row) for row in sheet.iter_rows(values_only=True)]}
                                for sheet in workbook.worksheets]
            result["formula_cells"] = []
            for sheet, cached in zip(formulas.worksheets, result["sheets"]):
                for row in sheet.iter_rows():
                    for cell in row:
                        if cell.data_type != "f":
                            continue
                        entry = {"sheet": sheet.title, "cell": cell.coordinate, "formula": cell.value,
                                 "cached_value": cached["rows"][cell.row - 1][cell.column - 1]}
                        result["formula_cells"].append(entry)
                        if entry["cached_value"] is None:
                            cached["rows"][cell.row - 1][cell.column - 1] = cell.value
                            result["errors"].append({"sheet": sheet.title, "cell": cell.coordinate,
                                                     "code": "formula_value_missing", "message": "Formula retained; no cached computed value"})
            result["complete"] = not result["errors"]
        finally:
            workbook.close()
            formulas.close()
        result["text"] = sheet_text(result["sheets"])
    elif extension == ".xls":
        import xlrd
        workbook = xlrd.open_workbook(str(path), on_demand=True)
        try:
            for sheet in workbook.sheets():
                if sheet.nrows * sheet.ncols > MAX_WORKSHEET_CELLS:
                    raise ValueError(f"Worksheet {sheet.name!r} has too many cells; explicitly reduce its dimensions before reading")
            result["sheets"] = [{"name": sheet.name, "rows": [sheet.row_values(i) for i in range(sheet.nrows)]}
                                for sheet in workbook.sheets()]
        finally:
            workbook.release_resources()
        result["text"] = sheet_text(result["sheets"])
    elif extension == ".pptx":
        from pptx import Presentation
        parts = []
        slides = Presentation(path).slides
        for number, slide in enumerate(slides, 1):
            parts.append(f"=== Слайд {number} ===")
            for shape in slide.shapes:
                if shape.has_text_frame:
                    parts.append(shape.text)
                if shape.has_table:
                    rows = [[cell.text for cell in row.cells] for row in shape.table.rows]
                    result["tables"].append({"slide": number, "rows": rows})
                    parts.extend(" | ".join(row) for row in rows)
        result["slide_count"] = len(slides)
        result["text"] = "\n".join(parts)
    elif extension == ".zip":
        parts, entries, total_bytes = [], [], 0
        with zipfile.ZipFile(path) as archive:
            if len(archive.infolist()) > 10000:
                raise ValueError("Archive has more than 10000 entries; explicitly select a smaller archive")
            for item in archive.infolist():
                entry = {"name": item.filename, "size": item.file_size, "text_read": False}
                entries.append(entry)
                parts.append(f"=== {item.filename} ({item.file_size} байт) ===")
                if not item.is_dir() and Path(item.filename).suffix.lower() in TEXT_EXTENSIONS:
                    if item.file_size > 1_000_000 or total_bytes + item.file_size > 20_000_000:
                        result["errors"].append({"entry": item.filename, "code": "entry_too_large"})
                        continue
                    try:
                        total_bytes += item.file_size
                        parts.append(decode_text(archive.read(item), encoding))
                        entry["text_read"] = True
                    except (ValueError, RuntimeError, OSError, zipfile.BadZipFile) as exc:
                        result["errors"].append({"entry": item.filename, "code": "entry_read_failed", "message": str(exc)})
        result["entries"] = entries
        result["complete"] = not result["errors"]
        result["text"] = "\n".join(parts)
    elif extension in TEXT_EXTENSIONS or not extension:
        result["text"] = decode_text(path.read_bytes(), encoding)
    else:
        raise ValueError(f"Unsupported format {extension}; binary DOC/PPT require explicit conversion to DOCX/PPTX")
    return {"ok": True, "source": str(path), "format": extension, **result}


def sheet_text(sheets: list[dict[str, Any]]) -> str:
    return "\n\n".join(f"=== Лист: {sheet['name']} ===\n" + "\n".join(
        " | ".join("" if value is None else str(value) for value in row) for row in sheet["rows"]) for sheet in sheets)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", type=Path, nargs="?")
    parser.add_argument("--url", help="Explicit HTTP(S) document URL, downloaded before parsing")
    parser.add_argument("--download", type=Path, help="Retain original URL bytes in this new file")
    parser.add_argument("--timeout", type=float, default=30)
    parser.add_argument("--max-bytes", type=int, default=25_000_000)
    parser.add_argument("--encoding", default="utf-8-sig", help="Explicit text/ZIP encoding; no ambiguous legacy guessing")
    parser.add_argument("--output", type=Path, help="Full UTF-8 text output")
    parser.add_argument("--json-output", type=Path, help="Full structured output")
    parser.add_argument("--offset", type=int, default=0)
    parser.add_argument("--max-chars", type=int, default=20000, help="stdout excerpt size; output files keep all text")
    args = parser.parse_args()
    try:
        if args.offset < 0 or args.max_chars < 1:
            raise ValueError("offset must be >= 0 and max-chars must be positive")
        if bool(args.input) == bool(args.url):
            raise ValueError("Provide exactly one input path or --url")
        if args.timeout <= 0 or args.max_bytes < 1:
            raise ValueError("timeout and max-bytes must be positive")
        provenance = None
        if args.url:
            if not args.download:
                raise ValueError("--url requires --download to retain original bytes")
            provenance = download_document(args.url, args.download, args.timeout, args.max_bytes)
            path = args.download.resolve(strict=True)
        else:
            path = args.input.expanduser().resolve(strict=True)
        if not path.is_file():
            raise ValueError("Input must be a file")
        for output in (args.output, args.json_output):
            if output and output.resolve() == path:
                raise ValueError("Output must not overwrite the input document")
        if args.output and args.json_output and args.output.resolve() == args.json_output.resolve():
            raise ValueError("Text and JSON outputs must use different paths")
        result = read_document(path, args.encoding)
        if provenance:
            result["source"] = provenance
        text = result["text"]
        if args.output:
            args.output.write_text(text, encoding="utf-8", newline="\n")
        if args.json_output:
            args.json_output.write_text(json.dumps(result, ensure_ascii=False, indent=2, default=str) + "\n", encoding="utf-8", newline="\n")
        end = args.offset + args.max_chars
        summary = {key: result[key] for key in ("ok", "source", "format", "complete", "ocr_required", "errors")}
        summary.update(text=text[args.offset:end], total_chars=len(text), offset=args.offset,
                       truncated=end < len(text), next_offset=end if end < len(text) else None,
                       page_count=len(result.get("pages", [])), table_count=len(result["tables"]),
                       output=str(args.output) if args.output else None,
                       json_output=str(args.json_output) if args.json_output else None)
        print(json.dumps(summary, ensure_ascii=False))
        return 0 if result["complete"] else 2
    except Exception as exc:
        print(json.dumps({"ok": False, "error": type(exc).__name__, "message": str(exc)}, ensure_ascii=False))
        return 1


if __name__ == "__main__":
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    raise SystemExit(main())
